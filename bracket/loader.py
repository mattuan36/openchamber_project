"""Workbook discovery, visual-bracket parsing, and validation."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .template import (
    CHAMPION_LABEL,
    NAME_PREFIX,
    REGION_LABELS,
    REGION_ORDER,
    ROUND_GAME_COUNTS,
    SEED_COLUMNS,
    SLOT_IDS,
    SLOTS_BY_ID,
    TEAM_COLUMNS,
    Slot,
    build_slots,
    normalize_team,
    parse_entrant_name,
    split_team_label,
)


class BracketError(Exception):
    """A user-facing input or template problem."""


SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".csv"}
DEFAULT_REPORT_NAME = "bracket_report.xlsx"
REPORT_TITLE = "2026 March Madness Bracket Report"


# ---------------------------------------------------------------------------
# Grid abstraction
# ---------------------------------------------------------------------------
class Grid:
    """A rectangular, merged-aware view of a worksheet or CSV export."""

    def __init__(self, values: list[list[str]], sheet_name: str, *, is_csv: bool = False):
        if not values:
            raise BracketError("The workbook or CSV is empty.")
        width = max(len(row) for row in values)
        self.values = [list(row) + [""] * (width - len(row)) for row in values]
        self.max_row = len(self.values)
        self.max_col = width
        self.sheet_name = sheet_name
        self.is_csv = is_csv

    def value(self, row: int, col: int) -> str:
        if row < 1 or col < 1 or row > self.max_row or col > self.max_col:
            return ""
        return self.values[row - 1][col - 1]


def _grid_from_worksheet(worksheet) -> Grid:
    max_row = max(worksheet.max_row, 1)
    max_col = max(worksheet.max_column, 1)
    values = [["" for _ in range(max_col)] for _ in range(max_row)]

    for row in worksheet.iter_rows(
        min_row=1, max_row=max_row, min_col=1, max_col=max_col
    ):
        for cell in row:
            value = cell.value
            if value is not None:
                values[cell.row - 1][cell.column - 1] = str(value).strip()

    # Merged cells hold their value in the top-left cell. Fill the whole range
    # so a mapped cell inside a merge still resolves correctly.
    for merged_range in worksheet.merged_cells.ranges:
        top_row = merged_range.min_row - 1
        top_col = merged_range.min_col - 1
        if top_row >= max_row or top_col >= max_col:
            continue
        merged_value = values[top_row][top_col]
        for row_index in range(
            merged_range.min_row - 1, min(merged_range.max_row, max_row)
        ):
            for col_index in range(
                merged_range.min_col - 1, min(merged_range.max_col, max_col)
            ):
                values[row_index][col_index] = merged_value

    return Grid(values, worksheet.title)


def _grid_from_csv(path: Path) -> Grid:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = [list(row) for row in csv.reader(handle)]
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise BracketError(f"Could not read CSV '{path}': {exc}") from exc
    if not rows:
        raise BracketError(f"The CSV '{path}' is empty.")
    width = max(len(row) for row in rows)
    values = [
        [cell.strip() for cell in row] + [""] * (width - len(row))
        for row in rows
    ]
    return Grid(values, path.name, is_csv=True)


def load_grid(path: Path) -> Grid:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _grid_from_csv(path)
    if suffix not in {".xlsx", ".xlsm"}:
        raise BracketError(
            f"Unsupported input '{path.name}'. Use .xlsx, .xlsm, or .csv for preview."
        )
    try:
        workbook = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:  # openpyxl raises a variety of parsing exceptions
        raise BracketError(f"Could not open workbook '{path}': {exc}") from exc

    try:
        valid: list[Grid] = []
        for worksheet in workbook.worksheets:
            grid = _grid_from_worksheet(worksheet)
            if _find_anchors(grid) is not None:
                valid.append(grid)
        if not valid:
            raise BracketError(
                f"Template mismatch in '{path.name}': no sheet contains the 2026 "
                "bracket anchors (NAME:, EAST/WEST/SOUTH/MIDWEST, round headers, "
                "and 2026 NCAA CHAMPION)."
            )
        if len(valid) == 1:
            return valid[0]
        active = workbook.active
        for grid in valid:
            if active is not None and grid.sheet_name == active.title:
                return grid
        names = ", ".join(grid.sheet_name for grid in valid)
        raise BracketError(
            f"Ambiguous workbook '{path.name}': multiple bracket sheets found ({names})."
        )
    finally:
        workbook.close()


def is_generated_report(path: Path) -> bool:
    """Identify a workbook written by this tool, so old reports are skipped."""

    if not path.exists() or path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return False
    try:
        workbook = load_workbook(path, read_only=True)
    except Exception:
        return False
    try:
        return (
            workbook.properties.title == REPORT_TITLE
            and (
                workbook.sheetnames == ["Picks", "Standings"]
                or workbook.sheetnames
                == ["Picks", "Standings", "State", "Teams", "Corrections"]
            )
        )
    finally:
        workbook.close()


# ---------------------------------------------------------------------------
# Anchors
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Anchors:
    name_row: int
    name_col: int
    east_row: int
    west_row: int
    south_row: int
    midwest_row: int
    champion_label_row: int
    champion_col: int


# Region labels are followed by a venue in parentheses or are the whole cell.
# Requiring that suffix avoids matching team names such as "South Florida".
_REGION_RE = re.compile(r"^(EAST|WEST|SOUTH|MIDWEST)(?:\s*\(|$)", re.IGNORECASE)
_ROUND_HEADERS = ("FIRST ROUND", "SECOND ROUND", "REGIONALS", "SEMIFINALS")


def _find_anchors(grid: Grid) -> Anchors | None:
    name_cell: tuple[int, int] | None = None
    champion_cell: tuple[int, int] | None = None
    regions: dict[str, tuple[int, int]] = {}
    observed_round_headers: set[str] = set()

    for row in range(1, grid.max_row + 1):
        for col in range(1, grid.max_col + 1):
            text = grid.value(row, col)
            if not text:
                continue
            collapsed = re.sub(r"\s+", " ", text).strip()
            upper = collapsed.upper()

            if upper.startswith(NAME_PREFIX.upper()) and name_cell is None:
                name_cell = (row, col)
            if upper == CHAMPION_LABEL.upper() and champion_cell is None:
                champion_cell = (row, col)

            for header in _ROUND_HEADERS:
                if upper.startswith(header):
                    observed_round_headers.add(header)

            match = _REGION_RE.match(collapsed)
            if not match:
                continue
            region_label = match.group(1).upper()
            region = next(
                (
                    key
                    for key, label in REGION_LABELS.items()
                    if label == region_label
                ),
                None,
            )
            if region is None or region in regions:
                continue
            # East/South labels sit in the left block; West/Midwest in the
            # right block. This avoids accidentally matching a team name.
            if region in {"E", "S"} and col > 3:
                continue
            if region in {"W", "M"} and col < 10:
                continue
            regions[region] = (row, col)

    if name_cell is None or champion_cell is None:
        return None
    if set(regions) != set(REGION_ORDER):
        return None
    if set(_ROUND_HEADERS) - observed_round_headers:
        return None

    east_row = regions["E"][0]
    west_row = regions["W"][0]
    south_row = regions["S"][0]
    midwest_row = regions["M"][0]
    champion_label_row, champion_col = champion_cell

    if east_row != west_row or south_row != midwest_row:
        return None
    if not (east_row < south_row < champion_label_row):
        return None
    # The label is normally in I, but a centered merge may start at H or J.
    # The champion pick itself remains I43 in the canonical map.
    if champion_col not in {8, 9, 10}:
        return None

    return Anchors(
        name_row=name_cell[0],
        name_col=name_cell[1],
        east_row=east_row,
        west_row=west_row,
        south_row=south_row,
        midwest_row=midwest_row,
        champion_label_row=champion_label_row,
        champion_col=champion_col,
    )


def _require_anchors(grid: Grid, path: Path) -> Anchors:
    anchors = _find_anchors(grid)
    if anchors is None:
        raise BracketError(
            f"Template mismatch in '{path.name}' ({grid.sheet_name}): expected "
            "NAME:, EAST/WEST/SOUTH/MIDWEST on matching rows, FIRST ROUND / "
            "SECOND ROUND / REGIONALS / SEMIFINALS, and 2026 NCAA CHAMPION in "
            "column I."
        )
    return anchors


# ---------------------------------------------------------------------------
# Bracket model
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Bracket:
    path: Path
    source_name: str
    entrant: str
    picks: dict[str, str]
    warnings: tuple[str, ...]
    sheet_name: str
    # The two first-round teams for each R64 slot, as printed on the bracket.
    # Needed for report generation and for validating edited picks.
    team_pairs: dict[str, tuple[str, ...]] = field(default_factory=dict)
    # Seeds aligned with team_pairs. Play-in teams repeat the shared seed.
    team_seeds: dict[str, tuple[int, ...]] = field(default_factory=dict)

    @property
    def filled(self) -> int:
        return sum(1 for value in self.picks.values() if value.strip())


@dataclass(frozen=True)
class Dataset:
    actual: Bracket
    entrants: list[Bracket]
    skipped: list[tuple[Path, str]]
    files: list[Path]
    warnings: tuple[str, ...]
    # Every supported file seen in the input location, including a report
    # path that discovery must ignore. The report writer uses this to avoid
    # overwriting an input that was excluded as the output.
    protected_paths: tuple[Path, ...] = ()

    @property
    def pending(self) -> int:
        return sum(1 for value in self.actual.picks.values() if not value.strip())


def _read_seeds(
    grid: Grid, slots: Iterable[Slot], path: Path
) -> dict[str, tuple[int, ...]]:
    """Read the two printed seeds for every R64 game.

    The returned tuples follow the printed row order. Validation still happens
    here so a malformed template fails before any bracket data is used.
    """

    slot_seeds: dict[str, tuple[int, ...]] = {}
    for region in REGION_ORDER:
        seeds: list[int] = []
        for slot in slots:
            if slot.round != "R64" or slot.region != region:
                continue
            row_seeds: list[int] = []
            for team_row, _team_col in slot.team_cells:
                raw_seed = grid.value(team_row, SEED_COLUMNS[region])
                try:
                    seed = int(float(raw_seed))
                except (TypeError, ValueError):
                    raise BracketError(
                        f"Template mismatch in '{path.name}': missing or invalid "
                        f"{region} seed at row {team_row}."
                    ) from None
                row_seeds.append(seed)
                seeds.append(seed)
            slot_seeds[slot.id] = tuple(row_seeds)
        if sorted(seeds) != list(range(1, 17)):
            raise BracketError(
                f"Template mismatch in '{path.name}': {region} seeds are "
                f"{sorted(seeds)}, expected 1-16."
            )
    return slot_seeds


def _detect_team_columns(grid: Grid, slots: Iterable[Slot]) -> dict[str, int]:
    """Resolve the right-hand team column from the populated R64 labels.

    The printable sheet mirrors the left team block; depending on the source
    workbook the right team name lands in P or Q (and a merged P:Q cell fills
    both). Left-hand team names are consistently in column A.
    """

    columns = dict(TEAM_COLUMNS)
    slot_list = list(slots)
    for region in ("W", "M"):
        scores: dict[int, int] = {}
        for candidate in (16, 17):
            scores[candidate] = 0
            for slot in slot_list:
                if slot.round != "R64" or slot.region != region:
                    continue
                for row, _col in slot.team_cells:
                    value = grid.value(row, candidate)
                    # A team label carries letters; a record such as "32-2"
                    # does not. This picks Q when P holds a record and P when
                    # the template's team block is one column to the left.
                    if len(re.findall(r"[A-Za-z]", value)) >= 2:
                        scores[candidate] += 1
        best = max((16, 17), key=lambda col: (scores[col], col == 17))
        if scores[best] == 0:
            raise BracketError(
                f"Template mismatch in '{grid.sheet_name}': no {region} team "
                "labels found in the right-hand team block."
            )
        columns[region] = best
    return columns


def _apply_team_columns(
    slots: Iterable[Slot], team_columns: dict[str, int]
) -> tuple[Slot, ...]:
    adjusted: list[Slot] = []
    for slot in slots:
        if slot.round == "R64" and slot.region in team_columns:
            slot = replace(
                slot,
                team_cells=tuple(
                    (row, team_columns[slot.region])
                    for row, _col in slot.team_cells
                ),
            )
        adjusted.append(slot)
    return tuple(adjusted)


def _read_team_pairs(
    grid: Grid,
    slots: Iterable[Slot],
    seed_rows: dict[str, tuple[int, ...]],
) -> tuple[dict[str, tuple[str, ...]], dict[str, tuple[int, ...]]]:
    pairs: dict[str, tuple[str, ...]] = {}
    seed_pairs: dict[str, tuple[int, ...]] = {}
    for slot in slots:
        if slot.round != "R64":
            continue
        teams: list[str] = []
        seeds: list[int] = []
        row_seeds = seed_rows.get(slot.id, ())
        for index, (row, col) in enumerate(slot.team_cells):
            label_teams = split_team_label(grid.value(row, col))
            if not label_teams:
                continue
            teams.extend(label_teams)
            # A play-in label such as "Texas / NC State" shares one printed
            # seed row, so every team split from that row carries that seed.
            if index < len(row_seeds):
                seeds.extend([row_seeds[index]] * len(label_teams))
        if not teams:
            raise BracketError(f"Template mismatch: no teams found for {slot.id}.")
        pairs[slot.id] = tuple(teams)
        if len(seeds) == len(teams):
            seed_pairs[slot.id] = tuple(seeds)
    return pairs, seed_pairs


def _validate_picks(
    picks: dict[str, str],
    slots: Iterable[Slot],
    team_pairs: dict[str, tuple[str, ...]],
) -> tuple[str, ...]:
    warnings: list[str] = []
    known = {
        normalize_team(team)
        for teams in team_pairs.values()
        for team in teams
        if normalize_team(team)
    }
    slot_list = list(slots)

    for slot in slot_list:
        pick = picks[slot.id].strip()
        if not pick:
            continue
        normalized = normalize_team(pick)
        if normalized not in known:
            warnings.append(f"{slot.id}: unrecognized team '{pick}'.")

        if slot.round == "R64":
            allowed = {normalize_team(team) for team in team_pairs[slot.id]}
            if normalized not in allowed:
                names = ", ".join(team_pairs[slot.id])
                warnings.append(
                    f"{slot.id}: pick '{pick}' is not in the first-round matchup "
                    f"({names})."
                )
            continue

        if not slot.feeders:
            continue
        feeder_picks = [picks.get(feeder, "").strip() for feeder in slot.feeders]
        if all(feeder_picks):
            allowed = {normalize_team(value) for value in feeder_picks}
            if normalized not in allowed:
                warnings.append(
                    f"{slot.id}: pick '{pick}' is not one of its feeder winners "
                    f"({', '.join(feeder_picks)})."
                )
    return tuple(warnings)


def parse_bracket(path: Path) -> Bracket:
    path = path.resolve()
    if not path.exists():
        raise BracketError(f"Input file not found: {path}")
    grid = load_grid(path)
    anchors = _require_anchors(grid, path)
    slots = build_slots(
        anchors.east_row,
        anchors.south_row,
        anchors.champion_label_row,
    )
    seed_rows = _read_seeds(grid, slots, path)
    slots = _apply_team_columns(slots, _detect_team_columns(grid, slots))
    team_pairs, team_seeds = _read_team_pairs(grid, slots, seed_rows)

    picks = {
        slot.id: grid.value(slot.row, slot.col).strip()
        for slot in slots
    }
    if set(picks) != set(SLOT_IDS):
        raise BracketError(
            f"Template mismatch in '{path.name}': the slot map did not produce "
            "all 63 games."
        )

    raw_name = grid.value(anchors.name_row, anchors.name_col)
    entrant = parse_entrant_name(raw_name, path.stem)
    warnings = list(_validate_picks(picks, slots, team_pairs))
    if grid.is_csv:
        warnings.append(
            "CSV preview: merged-cell information is unavailable; values were "
            "read from the CSV grid."
        )

    return Bracket(
        path=path,
        source_name=path.name,
        entrant=entrant,
        picks=picks,
        warnings=tuple(warnings),
        sheet_name=grid.sheet_name,
        team_pairs={slot_id: tuple(teams) for slot_id, teams in team_pairs.items()},
        team_seeds={
            slot_id: tuple(seeds) for slot_id, seeds in team_seeds.items()
        },
    )


# ---------------------------------------------------------------------------
# Discovery and actual-results selection
# ---------------------------------------------------------------------------
def _is_ignored(path: Path, output_path: Path | None) -> bool:
    if path.name.startswith("~$") or path.name.startswith("."):
        return True
    if output_path is not None and path.resolve() == output_path.resolve():
        return True
    return False


def discover_files(input_path: Path, output_path: Path | None = None) -> list[Path]:
    input_path = input_path.resolve()
    if not input_path.exists():
        raise BracketError(f"Input path not found: {input_path}")
    if input_path.is_file():
        if input_path.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise BracketError(
                f"Unsupported input '{input_path.name}'. Use .xlsx, .xlsm, or .csv."
            )
        return [input_path]
    if not input_path.is_dir():
        raise BracketError(f"Input path is not a file or directory: {input_path}")

    files = sorted(
        (
            path
            for path in input_path.iterdir()
            if path.is_file()
            and path.suffix.lower() in SUPPORTED_SUFFIXES
            and not _is_ignored(path, output_path)
        ),
        key=lambda path: (path.name.casefold(), path.name),
    )
    if not files:
        raise BracketError(
            f"No supported bracket files found in '{input_path}'. Expected "
            ".xlsx, .xlsm, or .csv."
        )
    return files


def select_actual_file(
    files: list[Path],
    actual_path: Path | None,
    output_path: Path | None,
) -> Path:
    if actual_path is not None:
        resolved = actual_path.resolve()
        if not resolved.exists() or not resolved.is_file():
            raise BracketError(f"--actual file not found: {resolved}")
        if resolved.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise BracketError(
                f"--actual must be .xlsx, .xlsm, or .csv, not '{resolved.suffix}'."
            )
        if output_path is not None and resolved == output_path.resolve():
            raise BracketError("The output path cannot also be the --actual file.")
        return resolved

    candidates = [
        path
        for path in files
        if "results" in path.stem.casefold()
        and (output_path is None or path.resolve() != output_path.resolve())
    ]
    if len(candidates) == 1:
        return candidates[0].resolve()
    if len(candidates) > 1:
        names = "\n".join(f"  - {path.name}" for path in candidates)
        raise BracketError(
            "Multiple files contain 'results' in their name. Pass --actual to "
            f"choose one:\n{names}"
        )
    raise BracketError(
        "No actual-results workbook found. Put 'results' in its filename "
        "(for example 2026_results.xlsx) or pass --actual PATH."
    )


def load_dataset(
    input_path: Path,
    actual_path: Path | None = None,
    output_path: Path | None = None,
    *,
    skip_invalid: bool = False,
) -> Dataset:
    output_path = output_path.resolve() if output_path is not None else None
    all_files = [
        path
        for path in discover_files(input_path, None)
        if not is_generated_report(path)
    ]
    files = [
        path
        for path in all_files
        if output_path is None or path.resolve() != output_path
    ]
    if not files:
        raise BracketError(
            "No supported bracket files remain after excluding the output path."
        )
    actual_file = select_actual_file(files, actual_path, output_path)

    actual = parse_bracket(actual_file)
    if actual.filled == 0:
        raise BracketError(
            f"Actual-results workbook '{actual.source_name}' contains no picks."
        )

    all_paths = [actual_file]
    for path in files:
        if path.resolve() != actual_file:
            all_paths.append(path)
    unique_paths = list(dict.fromkeys(path.resolve() for path in all_paths))

    entrant_paths = [path for path in unique_paths if path != actual_file]
    entrants: list[Bracket] = []
    skipped: list[tuple[Path, str]] = []
    for path in sorted(
        entrant_paths, key=lambda item: (item.name.casefold(), item.name)
    ):
        try:
            entrants.append(parse_bracket(path))
        except BracketError as exc:
            if skip_invalid:
                skipped.append((path, str(exc)))
                continue
            raise BracketError(f"{path.name}: {exc}") from exc

    entrants.sort(
        key=lambda bracket: (
            bracket.entrant.casefold(),
            bracket.source_name.casefold(),
        )
    )
    warnings: list[str] = []

    if not entrants:
        warnings.append(
            "No entrant brackets found; the report will contain the actual row only."
        )

    by_name: dict[str, list[tuple[str, str]]] = {}
    for bracket in entrants:
        by_name.setdefault(bracket.entrant.casefold(), []).append(
            (bracket.entrant, bracket.source_name)
        )
    for entries in by_name.values():
        if len(entries) > 1:
            display_name = entries[0][0]
            sources = ", ".join(source for _name, source in entries)
            warnings.append(
                f"Duplicate entrant name '{display_name}': {sources}. "
                "The Source File column disambiguates them."
            )

    for bracket in entrants:
        if bracket.filled == 0:
            warnings.append(
                f"'{bracket.source_name}' contains no picks; it will score 0."
            )

    # The actual bracket supplies the Teams sheet. Flag any entrant whose
    # printed first-round matchup disagrees, rather than silently choosing one.
    for bracket in entrants:
        for slot_id, actual_teams in actual.team_pairs.items():
            other_teams = bracket.team_pairs.get(slot_id)
            if other_teams is None:
                continue
            if tuple(map(normalize_team, other_teams)) != tuple(
                map(normalize_team, actual_teams)
            ):
                warnings.append(
                    f"{bracket.source_name}: {slot_id} first-round teams "
                    f"({', '.join(other_teams)}) differ from the actual bracket "
                    f"({', '.join(actual_teams)}). The actual bracket is used."
                )

    return Dataset(
        actual=actual,
        entrants=entrants,
        skipped=skipped,
        files=unique_paths,
        warnings=tuple(warnings),
        protected_paths=tuple(path.resolve() for path in all_files),
    )


def template_slot_counts() -> dict[str, int]:
    """Expose the expected per-round counts for reports and tests."""
    return dict(ROUND_GAME_COUNTS)


def slot_for_id(slot_id: str) -> Slot:
    try:
        return SLOTS_BY_ID[slot_id]
    except KeyError as exc:
        raise BracketError(f"Unknown slot '{slot_id}'.") from exc
