"""Formatted two-sheet Excel report for the bracket scorekeeper."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .loader import REPORT_TITLE, Dataset, is_generated_report
from .scoring import EntrantScore, TIE_BREAK_DESCRIPTION
from .template import (
    REGION_NAMES,
    ROUND_NAMES,
    ROUND_ORDER,
    ROUND_STANDINGS_HEADERS,
    ROUND_WEIGHTS,
    SLOT_IDS,
    SLOTS_BY_ID,
    TOTAL_MAX_POINTS,
    team_matches,
)


class ReportError(Exception):
    """A user-facing report-writing problem."""


# Palette: dark, round-coded headers with pale data fills.
ROUND_HEADER_FILLS = {
    "R64": "2F5597",
    "R32": "548235",
    "S16": "BF8F00",
    "E8": "C55A11",
    "F4": "7030A0",
    "CH": "C00000",
}
ROUND_DATA_FILLS = {
    "R64": "EAF1FB",
    "R32": "E2EFDA",
    "S16": "FFF2CC",
    "E8": "FCE4D6",
    "F4": "E4DFEC",
    "CH": "FFE699",
}
WHITE = "FFFFFF"
BLACK = "000000"
DARK_BLUE = "1F4E78"
GREEN_FILL = "C6EFCE"
GREEN_TEXT = "006100"
RED_FILL = "FFC7CE"
RED_TEXT = "9C0006"
PENDING_FILL = "E7E6E6"
PENDING_TEXT = "7F7F7F"
MUTED_TEXT = "595959"

HEADER_FONT = Font(name="Calibri", size=9, bold=True, color=WHITE)
BODY_FONT = Font(name="Calibri", size=9)
BODY_BOLD = Font(name="Calibri", size=9, bold=True)
NOTE_FONT = Font(name="Calibri", size=9, italic=True, color=MUTED_TEXT)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN_WHITE = Side(style="thin", color=WHITE)
THICK_BLACK = Side(style="thick", color=BLACK)
THIN_GRAY = Side(style="thin", color="D9D9D9")


def _slot_header(slot_id: str) -> str:
    slot = SLOTS_BY_ID[slot_id]
    if slot.round == "CH":
        return "Championship"
    if slot.round == "F4":
        return f"F4 Game {slot.game}"
    region = REGION_NAMES[slot.region or ""]
    return f"{slot.round} {region} {slot.game}"


def _slot_comment(slot_id: str) -> str:
    slot = SLOTS_BY_ID[slot_id]
    if slot.round == "CH":
        return f"Championship ({ROUND_WEIGHTS['CH']} points) at cell I{slot.row}."
    if slot.round == "F4":
        return f"Final Four game {slot.game} ({ROUND_WEIGHTS['F4']} points)."
    region = REGION_NAMES[slot.region or ""]
    return (
        f"{ROUND_NAMES[slot.round]} — {region} game {slot.game} "
        f"({ROUND_WEIGHTS[slot.round]} point"
        f"{'s' if ROUND_WEIGHTS[slot.round] != 1 else ''})."
    )


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _set_cell(
    cell,
    value,
    *,
    font=BODY_FONT,
    fill: str | None = None,
    alignment=CENTER,
    border: Border | None = None,
    number_format: str | None = None,
) -> None:
    cell.value = value
    cell.font = font
    cell.alignment = alignment
    if fill:
        cell.fill = _fill(fill)
    if border is not None:
        cell.border = border
    if number_format:
        cell.number_format = number_format


def _build_picks_sheet(
    workbook: Workbook, dataset: Dataset, scores: list[EntrantScore]
) -> None:
    sheet = workbook.active
    sheet.title = "Picks"
    sheet.sheet_properties.tabColor = DARK_BLUE
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 80

    headers = ["Source File", "Entrant", *(_slot_header(slot_id) for slot_id in SLOT_IDS)]
    for index, label in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        slot_id = SLOT_IDS[index - 3] if index >= 3 else None
        round_name = SLOTS_BY_ID[slot_id].round if slot_id else None
        header_fill = (
            ROUND_HEADER_FILLS[round_name] if round_name else DARK_BLUE
        )
        left_border = THIN_WHITE
        if index == 3:
            left_border = THICK_BLACK
        elif slot_id is not None:
            previous = SLOTS_BY_ID[SLOT_IDS[index - 4]]
            if previous.round != round_name:
                left_border = THICK_BLACK
            elif previous.region != SLOTS_BY_ID[slot_id].region:
                left_border = THIN_WHITE
        cell.font = HEADER_FONT
        cell.fill = _fill(header_fill)
        cell.alignment = CENTER
        cell.border = Border(left=left_border)
        if slot_id is not None:
            cell.comment = Comment(_slot_comment(slot_id), "Scorekeeper")
            cell.comment.width = 260
            cell.comment.height = 55

    sheet.freeze_panes = "C2"
    sheet.row_dimensions[1].height = 46
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 25
    for index in range(3, len(SLOT_IDS) + 3):
        sheet.column_dimensions[get_column_letter(index)].width = 14

    # The actual-results row comes first and is visually separate from entrants.
    actual_row = 2
    _set_cell(
        sheet.cell(row=actual_row, column=1),
        dataset.actual.source_name,
        font=Font(name="Calibri", size=9, bold=True, color=WHITE),
        fill=DARK_BLUE,
        alignment=LEFT_WRAP,
    )
    _set_cell(
        sheet.cell(row=actual_row, column=2),
        f"ACTUAL — {dataset.actual.entrant}",
        font=Font(name="Calibri", size=9, bold=True, color=WHITE),
        fill=DARK_BLUE,
        alignment=LEFT_WRAP,
    )
    for offset, slot_id in enumerate(SLOT_IDS, start=3):
        value = dataset.actual.picks.get(slot_id, "").strip()
        cell = sheet.cell(row=actual_row, column=offset)
        if value:
            _set_cell(
                cell,
                value,
                font=Font(name="Calibri", size=9, bold=True, color=WHITE),
                fill=DARK_BLUE,
            )
        else:
            _set_cell(
                cell,
                None,
                font=Font(name="Calibri", size=9, color=PENDING_TEXT),
                fill=PENDING_FILL,
            )

    # Keep the Picks rows in the same order as Standings (best rank first).
    entrant_brackets = (
        [score.bracket for score in scores] if scores else dataset.entrants
    )
    for row_index, bracket in enumerate(entrant_brackets, start=3):
        _set_cell(
            sheet.cell(row=row_index, column=1),
            bracket.source_name,
            alignment=LEFT_WRAP,
        )
        _set_cell(
            sheet.cell(row=row_index, column=2),
            bracket.entrant,
            font=BODY_BOLD,
            alignment=LEFT_WRAP,
        )
        for offset, slot_id in enumerate(SLOT_IDS, start=3):
            actual_value = dataset.actual.picks.get(slot_id, "").strip()
            pick = bracket.picks.get(slot_id, "").strip()
            cell = sheet.cell(row=row_index, column=offset)
            if not actual_value:
                _set_cell(
                    cell,
                    pick or None,
                    font=Font(name="Calibri", size=9, color=PENDING_TEXT),
                    fill=PENDING_FILL,
                )
            elif pick and team_matches(pick, actual_value):
                _set_cell(
                    cell,
                    pick,
                    font=Font(name="Calibri", size=9, color=GREEN_TEXT),
                    fill=GREEN_FILL,
                )
            else:
                _set_cell(
                    cell,
                    pick or None,
                    font=Font(name="Calibri", size=9, color=RED_TEXT),
                    fill=RED_FILL,
                )
        sheet.row_dimensions[row_index].height = 30

    sheet.row_dimensions[actual_row].height = 30
    last_row = max(2, 2 + len(dataset.entrants))
    last_col = 2 + len(SLOT_IDS)
    sheet.auto_filter.ref = f"A1:{get_column_letter(last_col)}{last_row}"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True

    note = (
        "Row 2 is the actual-results bracket. Green = correct, "
        "red = incorrect, gray = pending (the actual cell is blank)."
    )
    sheet["A1"].comment = Comment(note, "Scorekeeper")
    sheet["A1"].comment.width = 320
    sheet["A1"].comment.height = 80


def _build_standings_sheet(
    workbook: Workbook, dataset: Dataset, scores: list[EntrantScore]
) -> None:
    sheet = workbook.create_sheet("Standings")
    sheet.sheet_properties.tabColor = "548235"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90

    headers = [
        "Rank",
        "Entrant",
        "Source File",
        *[ROUND_STANDINGS_HEADERS[round_name] for round_name in ROUND_ORDER],
        "Total",
        "Correct Picks",
        "Pending Games",
        "Accuracy",
    ]
    for index, label in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.font = HEADER_FONT
        cell.fill = _fill(DARK_BLUE)
        cell.alignment = CENTER
        cell.border = Border(bottom=THIN_GRAY)
    sheet.row_dimensions[1].height = 44

    widths = {
        "A": 7,
        "B": 27,
        "C": 36,
        "D": 11,
        "E": 11,
        "F": 11,
        "G": 12,
        "H": 11,
        "I": 17,
        "J": 10,
        "K": 12,
        "L": 12,
        "M": 11,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    if not scores:
        sheet.merge_cells("A2:M2")
        cell = sheet["A2"]
        _set_cell(
            cell,
            "No entrant brackets were found. Add filled bracket workbooks to "
            "the input folder and run again.",
            font=NOTE_FONT,
            alignment=LEFT_WRAP,
        )
        sheet.row_dimensions[2].height = 32
    else:
        for row_index, score in enumerate(scores, start=2):
            rank_fill = {1: "FFD966", 2: "D9E1F2", 3: "F4B183"}.get(score.rank)
            rank_font = BODY_BOLD if score.rank <= 3 else BODY_FONT
            _set_cell(
                sheet.cell(row=row_index, column=1),
                score.rank,
                font=rank_font,
                fill=rank_fill,
            )
            _set_cell(
                sheet.cell(row=row_index, column=2),
                score.entrant,
                font=rank_font,
                alignment=LEFT_WRAP,
            )
            _set_cell(
                sheet.cell(row=row_index, column=3),
                score.source_name,
                alignment=LEFT_WRAP,
            )
            for column, round_name in enumerate(ROUND_ORDER, start=4):
                _set_cell(
                    sheet.cell(row=row_index, column=column),
                    score.round_scores[round_name],
                )
            _set_cell(
                sheet.cell(row=row_index, column=10),
                score.total,
                font=BODY_BOLD,
            )
            _set_cell(
                sheet.cell(row=row_index, column=11),
                score.correct,
            )
            _set_cell(
                sheet.cell(row=row_index, column=12),
                score.pending,
            )
            _set_cell(
                sheet.cell(row=row_index, column=13),
                score.accuracy,
                number_format="0.0%",
            )
            sheet.row_dimensions[row_index].height = 22

        last_row = 1 + len(scores)
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = f"A1:M{last_row}"
        sheet.conditional_formatting.add(
            f"J2:J{last_row}",
            DataBarRule(
                start_type="num",
                start_value=0,
                end_type="num",
                end_value=TOTAL_MAX_POINTS,
                color="5B9BD5",
                showValue=True,
            ),
        )

        note_row = last_row + 2
        scoring_note = (
            "Scoring: Round of 64 = 1, Round of 32 = 2, Sweet 16 = 4, "
            "Elite Eight = 6, Final Four = 8, Championship = 12. "
            f"Maximum {TOTAL_MAX_POINTS} points."
        )
        pending_note = (
            "A blank actual result is pending: it awards no points and is "
            "excluded from accuracy. Accuracy = correct picks ÷ decided games "
            "(not ÷ 63)."
        )
        for row_offset, text in enumerate(
            (scoring_note, pending_note, TIE_BREAK_DESCRIPTION)
        ):
            row = note_row + row_offset
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)
            _set_cell(
                sheet.cell(row=row, column=1),
                text,
                font=NOTE_FONT,
                alignment=LEFT_WRAP,
            )
            sheet.row_dimensions[row].height = 30

    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True

    sheet["J1"].comment = Comment(
        f"Total points out of {TOTAL_MAX_POINTS}; the bar is scaled to the maximum.",
        "Scorekeeper",
    )
    sheet["M1"].comment = Comment(
        "Correct picks divided by decided games. Pending games are excluded.",
        "Scorekeeper",
    )


def build_report(dataset: Dataset, scores: list[EntrantScore]) -> Workbook:
    """Build the report workbook in memory."""

    workbook = Workbook()
    _build_picks_sheet(workbook, dataset, scores)
    _build_standings_sheet(workbook, dataset, scores)
    workbook.properties.title = REPORT_TITLE
    workbook.properties.creator = "Bracket Scorekeeper"
    workbook.properties.description = (
        "Actual results, entrant picks, standings, pending games, and accuracy."
    )
    return workbook


def write_report(
    dataset: Dataset,
    scores: list[EntrantScore],
    output_path: Path,
) -> Path:
    """Write the report atomically, refusing to overwrite an input file."""

    output_path = Path(output_path).resolve()
    input_paths = {path.resolve() for path in dataset.files}
    protected_paths = {path.resolve() for path in dataset.protected_paths}
    if output_path in input_paths or (
        output_path in protected_paths and not is_generated_report(output_path)
    ):
        raise ReportError(
            f"Refusing to overwrite an input file: {output_path}. Choose a "
            "different --output path."
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_report(dataset, scores)
    temp_name: str | None = None
    try:
        file_descriptor, temp_name = tempfile.mkstemp(
            prefix=f".{output_path.stem}_",
            suffix=".xlsx",
            dir=output_path.parent,
        )
        os.close(file_descriptor)
        workbook.save(temp_name)
        try:
            os.replace(temp_name, output_path)
        except PermissionError as exc:
            raise ReportError(
                f"Cannot replace '{output_path}'. Close it in Excel and try again."
            ) from exc
        temp_name = None
    except ReportError:
        raise
    except Exception as exc:
        raise ReportError(f"Could not write report '{output_path}': {exc}") from exc
    finally:
        workbook.close()
        if temp_name is not None:
            try:
                os.unlink(temp_name)
            except OSError:
                pass

    return output_path
