"""Persistence for the five-sheet report and the local-folder workflow.

The report is the source of truth after it is created. Raw bracket files are
used only to bootstrap or rebuild a report. This module keeps that policy in
one place so the CLI and the Streamlit app behave the same way.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field, replace
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from .loader import (
    DEFAULT_REPORT_NAME,
    Bracket,
    BracketError,
    Dataset,
    discover_files,
    is_generated_report,
    load_dataset,
)
from .report import (
    CORRECTION_FIELDS,
    TEAM_FIELDS,
    TEAM_SEED_FIELDS,
    ReportError,
    default_state,
    utc_now_iso,
    write_report,
)
from .scoring import EntrantScore, score_all
from .template import (
    HEADER_TO_SLOT,
    REGION_NAMES,
    REGION_ORDER,
    SLOT_IDS,
    SLOTS_BY_ID,
    normalize_team,
    split_team_label,
    team_matches,
)


class ReportLoadError(ReportError):
    """A saved report is missing columns or has an unusable shape."""


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _clean_header(value: object) -> str:
    return re.sub(r"\s+", " ", _text(value)).strip()


def _header_maps() -> tuple[dict[str, str], dict[str, str]]:
    by_label = {label.casefold(): slot_id for label, slot_id in HEADER_TO_SLOT.items()}
    by_id = {slot_id.casefold(): slot_id for slot_id in SLOT_IDS}
    return by_label, by_id


def _picks_columns(sheet) -> tuple[int, int, dict[str, int]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ReportLoadError("The Picks sheet is empty.")

    by_label, by_id = _header_maps()
    source_col: int | None = None
    entrant_col: int | None = None
    slot_columns: dict[str, int] = {}
    for index, value in enumerate(rows[0]):
        header = _clean_header(value)
        folded = header.casefold()
        if folded == "source file":
            source_col = index
            continue
        if folded == "entrant":
            entrant_col = index
            continue
        slot_id = by_label.get(folded) or by_id.get(folded)
        if slot_id is None:
            continue
        if slot_id in slot_columns:
            raise ReportLoadError(
                f"The Picks sheet has more than one column for {slot_id}."
            )
        slot_columns[slot_id] = index

    missing = [slot_id for slot_id in SLOT_IDS if slot_id not in slot_columns]
    if missing:
        raise ReportLoadError(
            "The Picks sheet is missing game columns for: " + ", ".join(missing)
        )
    if source_col is None or entrant_col is None:
        raise ReportLoadError(
            "The Picks sheet must contain 'Source File' and 'Entrant' headers."
        )
    return source_col, entrant_col, slot_columns


def _pick_row(row: tuple, slot_columns: dict[str, int]) -> dict[str, str]:
    return {
        slot_id: _text(row[column]) if column < len(row) else ""
        for slot_id, column in slot_columns.items()
    }


def _validate_loaded_picks(
    label: str,
    picks: dict[str, str],
    team_pairs: dict[str, tuple[str, ...]],
) -> list[str]:
    if not team_pairs:
        return []
    warnings: list[str] = []
    known = {
        normalize_team(team)
        for teams in team_pairs.values()
        for team in teams
        if normalize_team(team)
    }
    for slot_id in SLOT_IDS:
        pick = picks.get(slot_id, "").strip()
        if not pick:
            continue
        normalized = normalize_team(pick)
        if normalized not in known:
            warnings.append(f"{label}: {slot_id}: unrecognized team '{pick}'.")
            continue
        slot = SLOTS_BY_ID[slot_id]
        if slot.round == "R64":
            allowed = {normalize_team(team) for team in team_pairs.get(slot_id, ())}
            if allowed and normalized not in allowed:
                warnings.append(
                    f"{label}: {slot_id}: pick '{pick}' is not in its "
                    "first-round matchup."
                )
            continue
        feeder_picks = [picks.get(feeder, "").strip() for feeder in slot.feeders]
        if all(feeder_picks) and not any(
            team_matches(pick, value) for value in feeder_picks
        ):
            warnings.append(
                f"{label}: {slot_id}: pick '{pick}' is not one of its feeder "
                f"winners ({', '.join(feeder_picks)})."
            )
    return warnings


def _load_seed(value: object, slot_id: str, field_name: str) -> int:
    try:
        seed = int(float(_text(value)))
    except (TypeError, ValueError):
        raise ReportLoadError(
            f"Teams has a missing or invalid {field_name} for {slot_id}."
        ) from None
    if seed < 1 or seed > 16:
        raise ReportLoadError(
            f"Teams has an invalid {field_name} for {slot_id}: {seed}."
        )
    return seed


def _load_teams(
    sheet,
) -> tuple[dict[str, tuple[str, ...]], dict[str, tuple[int, ...]]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ReportLoadError("The Teams sheet is empty.")
    indices: dict[str, int] = {}
    for index, value in enumerate(rows[0]):
        header = _clean_header(value).casefold()
        for field_name in (*TEAM_FIELDS, *TEAM_SEED_FIELDS):
            if header == field_name.casefold():
                indices[field_name] = index
    missing = [field_name for field_name in TEAM_FIELDS if field_name not in indices]
    if missing:
        raise ReportLoadError(
            "The Teams sheet is missing headers: " + ", ".join(missing)
        )
    present_seed_fields = [
        field_name for field_name in TEAM_SEED_FIELDS if field_name in indices
    ]
    if present_seed_fields and len(present_seed_fields) != len(TEAM_SEED_FIELDS):
        raise ReportLoadError(
            "The Teams sheet must contain both Seed 1 and Seed 2 headers."
        )
    seeded = len(present_seed_fields) == len(TEAM_SEED_FIELDS)

    pairs: dict[str, tuple[str, ...]] = {}
    team_seeds: dict[str, tuple[int, ...]] = {}
    region_seed_cells: dict[str, list[int]] = {
        region: [] for region in REGION_ORDER
    }
    for row in rows[1:]:
        slot_id = _text(row[indices["Slot"]])
        if not slot_id:
            continue
        if slot_id not in SLOTS_BY_ID or not slot_id.startswith("R64-"):
            raise ReportLoadError(f"Teams contains an unknown slot: {slot_id}")
        if slot_id in pairs:
            raise ReportLoadError(f"Teams contains duplicate slot: {slot_id}")
        team_one = _text(row[indices["Team 1"]])
        team_two = _text(row[indices["Team 2"]])
        if not team_one or not team_two:
            raise ReportLoadError(f"Teams is missing a team for {slot_id}.")
        row_one = tuple(split_team_label(team_one))
        row_two = tuple(split_team_label(team_two))
        if not row_one or not row_two:
            raise ReportLoadError(f"Teams is missing a team for {slot_id}.")
        pairs[slot_id] = (*row_one, *row_two)
        if seeded:
            seed_one = _load_seed(
                row[indices["Seed 1"]], slot_id, "Seed 1"
            )
            seed_two = _load_seed(
                row[indices["Seed 2"]], slot_id, "Seed 2"
            )
            team_seeds[slot_id] = (
                *([seed_one] * len(row_one)),
                *([seed_two] * len(row_two)),
            )
            region = SLOTS_BY_ID[slot_id].region or ""
            region_seed_cells.setdefault(region, []).extend((seed_one, seed_two))

    missing_slots = [
        slot_id
        for slot_id in SLOT_IDS
        if slot_id.startswith("R64-") and slot_id not in pairs
    ]
    if missing_slots:
        raise ReportLoadError(
            "Teams is missing first-round matchups: " + ", ".join(missing_slots)
        )
    if seeded:
        for region in REGION_ORDER:
            if sorted(region_seed_cells.get(region, ())) != list(range(1, 17)):
                raise ReportLoadError(
                    f"Teams {REGION_NAMES[region]} seed set is "
                    f"{sorted(region_seed_cells.get(region, ()))}, expected 1-16."
                )
    return pairs, team_seeds


def _load_state(sheet) -> dict[str, str]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {}
    key_col = value_col = None
    for index, value in enumerate(rows[0]):
        header = _clean_header(value).casefold()
        if header == "key":
            key_col = index
        elif header == "value":
            value_col = index
    if key_col is None or value_col is None:
        raise ReportLoadError("The State sheet must contain 'Key' and 'Value' headers.")
    state: dict[str, str] = {}
    for row in rows[1:]:
        key = _text(row[key_col]) if key_col < len(row) else ""
        if not key:
            continue
        state[key] = _text(row[value_col]) if value_col < len(row) else ""
    return state


def _load_corrections(sheet) -> list[dict[str, str]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    indices: dict[str, int] = {}
    for index, value in enumerate(rows[0]):
        header = _clean_header(value)
        for field_name in CORRECTION_FIELDS:
            if header.casefold() == field_name.casefold():
                indices[field_name] = index
    missing = [
        field_name for field_name in CORRECTION_FIELDS if field_name not in indices
    ]
    if missing:
        raise ReportLoadError(
            "The Corrections sheet is missing headers: " + ", ".join(missing)
        )
    entries: list[dict[str, str]] = []
    for row in rows[1:]:
        if not any(_text(value) for value in row):
            continue
        entry = {
            field_name: (
                _text(row[column]) if column < len(row) else ""
            )
            for field_name, column in indices.items()
        }
        if isinstance(row[indices["Timestamp"]], datetime):
            timestamp = row[indices["Timestamp"]].isoformat()
            entry["Timestamp"] = timestamp
        entries.append(entry)
    return entries


def _loaded_state_defaults(path: Path) -> dict[str, str]:
    try:
        modified = datetime.fromtimestamp(
            path.stat().st_mtime, timezone.utc
        ).replace(microsecond=0).isoformat()
    except OSError:
        modified = utc_now_iso()
    return {
        "Report schema version": "1",
        "Created (UTC)": modified,
        "Updated (UTC)": modified,
        "Source folder": str(path.parent),
        "Scoring rules version": "legacy",
        "App version": "legacy report",
        "Report format": "Legacy two-sheet report",
    }


def load_report(path: Path) -> "ReportSession":
    """Load a five-sheet report, or a legacy Picks/Standings report."""

    path = Path(path).resolve()
    if not path.exists():
        raise ReportLoadError(f"Report not found: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ReportLoadError("Saved reports must be .xlsx or .xlsm workbooks.")
    try:
        workbook = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:
        raise ReportLoadError(f"Could not open report '{path}': {exc}") from exc

    try:
        names = workbook.sheetnames
        if "Picks" not in names or "Standings" not in names:
            raise ReportLoadError(
                "This workbook is not a scorekeeper report: it has no Picks "
                "and Standings sheets."
            )
        picks = workbook["Picks"]
        source_col, entrant_col, slot_columns = _picks_columns(picks)
        rows = list(picks.iter_rows(values_only=True))

        actual_index: int | None = None
        actual_name = ""
        for index, row in enumerate(rows[1:], start=1):
            label = _text(row[entrant_col]) if entrant_col < len(row) else ""
            match = re.match(r"^\s*ACTUAL\s*[—–-]\s*(.*)$", label, re.IGNORECASE)
            if match:
                actual_index = index
                actual_name = match.group(1).strip()
                break
        if actual_index is None:
            raise ReportLoadError(
                "Picks row 2 must be the actual-results row labelled "
                "'ACTUAL — <name>'."
            )

        teams: dict[str, tuple[str, ...]] = {}
        team_seeds: dict[str, tuple[int, ...]] = {}
        if "Teams" in names:
            teams, team_seeds = _load_teams(workbook["Teams"])

        actual_row = rows[actual_index]
        actual_source = (
            _text(actual_row[source_col]) if source_col < len(actual_row) else ""
        )
        actual_picks = _pick_row(actual_row, slot_columns)
        warnings: list[str] = []
        if not _complete_teams(teams):
            warnings.append(
                "Legacy report: the Teams sheet is missing first-round team data. "
                "Rebuild from the original bracket files before editing or saving."
            )
        warnings.extend(
            _validate_loaded_picks("ACTUAL", actual_picks, teams)
        )
        actual = Bracket(
            path=path,
            source_name=actual_source or path.name,
            entrant=actual_name or "ACTUAL",
            picks=actual_picks,
            warnings=tuple(
                _validate_loaded_picks(actual_name or "ACTUAL", actual_picks, teams)
            ),
            sheet_name="Picks",
            team_pairs=teams,
            team_seeds=team_seeds,
        )

        entrants: list[Bracket] = []
        for row in rows[actual_index + 1 :]:
            name = _text(row[entrant_col]) if entrant_col < len(row) else ""
            if not name or name.upper().startswith("ACTUAL"):
                continue
            source_name = _text(row[source_col]) if source_col < len(row) else ""
            picks_for_entrant = _pick_row(row, slot_columns)
            entrant_warnings = _validate_loaded_picks(name, picks_for_entrant, teams)
            warnings.extend(entrant_warnings)
            entrants.append(
                Bracket(
                    path=path,
                    source_name=source_name or f"{name}.xlsx",
                    entrant=name,
                    picks=picks_for_entrant,
                    warnings=tuple(entrant_warnings),
                    sheet_name="Picks",
                    team_pairs=teams,
                    team_seeds=team_seeds,
                )
            )

        state = _loaded_state_defaults(path)
        if "State" in names:
            state.update(_load_state(workbook["State"]))
        corrections: list[dict[str, str]] = []
        if "Corrections" in names:
            corrections = _load_corrections(workbook["Corrections"])

        dataset = Dataset(
            actual=actual,
            entrants=entrants,
            skipped=[],
            files=[],
            warnings=tuple(warnings),
            protected_paths=(),
        )
        return ReportSession(
            report_path=path,
            actual=actual,
            entrants=entrants,
            team_pairs=teams,
            team_seeds=team_seeds,
            state=state,
            corrections=corrections,
            source_folder=state.get("Source folder", str(path.parent)),
            warnings=tuple(warnings),
            _dataset=dataset,
        )
    finally:
        workbook.close()


@dataclass
class FolderScan:
    folder: Path
    raw_files: list[Path]
    reports: list[Path]
    actual_candidates: list[Path]
    default_report: Path


def scan_folder(folder: Path) -> FolderScan:
    """Scan one folder non-recursively, separating raw brackets from reports."""

    folder = Path(folder).resolve()
    files = discover_files(folder)
    reports = [path for path in files if is_generated_report(path)]
    raw_files = [path for path in files if path not in reports]
    return FolderScan(
        folder=folder,
        raw_files=raw_files,
        reports=reports,
        actual_candidates=[
            path for path in raw_files if "results" in path.stem.casefold()
        ],
        default_report=folder / DEFAULT_REPORT_NAME,
    )


def _validate_report_path(path: Path) -> Path:
    path = Path(path).resolve()
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ReportError("The report must be an .xlsx or .xlsm workbook.")
    if "results" in path.stem.casefold():
        raise ReportError(
            "The report filename must not contain 'results'; that word "
            "identifies the actual-results input during discovery."
        )
    return path


def _known_team_names(team_pairs: dict[str, tuple[str, ...]]) -> set[str]:
    return {
        normalize_team(team)
        for teams in team_pairs.values()
        for team in teams
        if normalize_team(team)
    }


def _complete_teams(team_pairs: dict[str, tuple[str, ...]]) -> bool:
    """True when every first-round matchup has both printed teams."""

    return all(
        len(team_pairs.get(slot_id, ())) >= 2
        for slot_id in SLOT_IDS
        if slot_id.startswith("R64-")
    )


def _complete_seeds(
    team_pairs: dict[str, tuple[str, ...]],
    team_seeds: dict[str, tuple[int, ...]],
) -> bool:
    """True when every expanded first-round team has a matching seed."""

    return bool(team_seeds) and all(
        len(team_pairs.get(slot_id, ())) >= 2
        and len(team_seeds.get(slot_id, ())) == len(team_pairs[slot_id])
        for slot_id in SLOT_IDS
        if slot_id.startswith("R64-")
    )


@dataclass
class ReportSession:
    """The editable state behind a saved report."""

    report_path: Path
    actual: Bracket
    entrants: list[Bracket]
    team_pairs: dict[str, tuple[str, ...]]
    team_seeds: dict[str, tuple[int, ...]]
    state: dict[str, str]
    corrections: list[dict[str, str]] = field(default_factory=list)
    source_folder: str = ""
    warnings: tuple[str, ...] = ()
    _dataset: Dataset | None = field(default=None, repr=False)

    @property
    def dataset(self) -> Dataset:
        if self._dataset is not None:
            return self._dataset
        return Dataset(
            actual=self.actual,
            entrants=self.entrants,
            skipped=[],
            files=[],
            warnings=self.warnings,
            protected_paths=(),
        )

    @property
    def scores(self) -> list[EntrantScore]:
        return score_all(self.actual, self.entrants)

    @property
    def pending(self) -> int:
        return sum(1 for value in self.actual.picks.values() if not value.strip())

    @property
    def teams_complete(self) -> bool:
        """True when every first-round matchup has both printed teams."""

        return _complete_teams(self.team_pairs)

    @property
    def team_seeds_complete(self) -> bool:
        """True when every expanded first-round team has a saved seed."""

        return _complete_seeds(self.team_pairs, self.team_seeds)

    @property
    def team_seed_map(self) -> dict[str, int]:
        """Normalized team name to its first-round seed."""

        seed_map: dict[str, int] = {}
        for slot_id, teams in self.team_pairs.items():
            for team, seed in zip(teams, self.team_seeds.get(slot_id, ())):
                normalized = normalize_team(team)
                if normalized:
                    seed_map.setdefault(normalized, seed)
        return seed_map

    @classmethod
    def from_dataset(
        cls,
        dataset: Dataset,
        report_path: Path,
        *,
        source_folder: str | Path | None = None,
        state: dict[str, str] | None = None,
        corrections: list[dict[str, str]] | None = None,
    ) -> "ReportSession":
        report_path = _validate_report_path(Path(report_path))
        current = dict(state or {})
        base = default_state(
            dataset,
            source_folder=source_folder,
            created=current.get("Created (UTC)"),
            updated=current.get("Updated (UTC)"),
        )
        base.update(current)
        base["Source folder"] = str(
            source_folder if source_folder is not None else base["Source folder"]
        )
        return cls(
            report_path=report_path,
            actual=dataset.actual,
            entrants=list(dataset.entrants),
            team_pairs=dict(dataset.actual.team_pairs),
            team_seeds=dict(dataset.actual.team_seeds),
            state=base,
            corrections=list(corrections or []),
            source_folder=str(
                source_folder if source_folder is not None else base["Source folder"]
            ),
            warnings=tuple(dataset.warnings),
        )

    def refreshed(self) -> "ReportSession":
        """Clear the cached dataset after pick changes."""

        self._dataset = None
        return self

    def save(self, path: Path | None = None) -> Path:
        if not self.teams_complete:
            raise ReportError(
                "This legacy report has no complete first-round team data. "
                "Rebuild it from the original bracket files before saving."
            )
        target = _validate_report_path(Path(path) if path is not None else self.report_path)
        if (
            target.exists()
            and target != self.report_path.resolve()
            and not is_generated_report(target)
        ):
            raise ReportError(
                f"Refusing to overwrite a raw bracket or non-report file: "
                f"{target}. Choose the saved report path or a new .xlsx name."
            )
        state = dict(self.state)
        state.update(default_state(self.dataset, source_folder=self.source_folder))
        state["Created (UTC)"] = self.state.get(
            "Created (UTC)", state["Created (UTC)"]
        )
        write_report(
            self.dataset,
            self.scores,
            target,
            state=state,
            corrections=self.corrections,
        )
        self.report_path = target
        self.state = state
        self._dataset = None
        return target

    def with_actual_pick(
        self, slot_id: str, pick: str
    ) -> tuple["ReportSession", list[str]]:
        if slot_id not in SLOTS_BY_ID:
            raise ReportError(f"Unknown slot '{slot_id}'.")
        picks = dict(self.actual.picks)
        picks[slot_id] = pick.strip()
        picks, cleared = _cascade(picks, slot_id, self.team_pairs)
        actual = replace(
            self.actual,
            picks=picks,
            team_pairs=self.team_pairs,
            team_seeds=self.team_seeds,
        )
        updated = replace(self, actual=actual, _dataset=None)
        return updated, cleared

    def with_corrected_entrant(
        self,
        entrant_index: int,
        slot_id: str,
        new_pick: str,
        reason: str,
        *,
        timestamp: str | None = None,
    ) -> tuple["ReportSession", dict[str, str]]:
        if not 0 <= entrant_index < len(self.entrants):
            raise ReportError("Choose a valid entrant to correct.")
        if slot_id not in SLOTS_BY_ID:
            raise ReportError(f"Unknown slot '{slot_id}'.")
        if not reason.strip():
            raise ReportError("A correction reason is required.")

        entrant = self.entrants[entrant_index]
        old_pick = entrant.picks.get(slot_id, "").strip()
        new_pick = new_pick.strip()
        if (old_pick and new_pick and team_matches(old_pick, new_pick)) or (
            not old_pick and not new_pick
        ):
            raise ReportError("The corrected pick is unchanged.")
        if new_pick and self.team_pairs:
            normalized = normalize_team(new_pick)
            known = _known_team_names(self.team_pairs)
            if normalized not in known:
                raise ReportError(f"'{new_pick}' is not a team in this bracket.")
            slot = SLOTS_BY_ID[slot_id]
            if slot.round == "R64":
                allowed = {
                    normalize_team(team)
                    for team in self.team_pairs.get(slot_id, ())
                }
                if normalized not in allowed:
                    raise ReportError(
                        f"'{new_pick}' is not in the {slot_id} first-round matchup."
                    )
            else:
                feeder_picks = [
                    entrant.picks.get(feeder, "").strip()
                    for feeder in slot.feeders
                ]
                if all(feeder_picks) and not any(
                    team_matches(new_pick, value) for value in feeder_picks
                ):
                    raise ReportError(
                        f"'{new_pick}' is not one of the feeder winners for "
                        f"{slot_id}."
                    )

        entry = {
            "Timestamp": timestamp or utc_now_iso(),
            "Entrant": entrant.entrant,
            "Source File": entrant.source_name,
            "Slot": slot_id,
            "Old Pick": old_pick,
            "New Pick": new_pick,
            "Reason": reason.strip(),
        }
        picks = dict(entrant.picks)
        picks[slot_id] = new_pick
        corrected = replace(
            entrant,
            picks=picks,
            team_pairs=self.team_pairs,
            team_seeds=self.team_seeds,
        )
        entrants = list(self.entrants)
        entrants[entrant_index] = corrected
        corrections = [*self.corrections, entry]
        updated = replace(
            self,
            entrants=entrants,
            corrections=corrections,
            _dataset=None,
        )
        return updated, entry


def _cascade(
    picks: dict[str, str],
    changed_slot: str,
    team_pairs: dict[str, tuple[str, ...]],
) -> tuple[dict[str, str], list[str]]:
    from .scoring import cascade_clear

    return cascade_clear(picks, changed_slot, team_pairs)


def build_from_raw(
    folder: Path,
    *,
    actual_path: Path | None = None,
    report_path: Path | None = None,
    skip_invalid: bool = True,
) -> ReportSession:
    """Build an in-memory report from the raw bracket files in a folder."""

    folder = Path(folder).resolve()
    scan = scan_folder(folder)
    target = _validate_report_path(
        Path(report_path) if report_path is not None else scan.default_report
    )
    dataset = load_dataset(
        folder,
        actual_path=Path(actual_path) if actual_path is not None else None,
        output_path=target,
        skip_invalid=skip_invalid,
    )
    return ReportSession.from_dataset(
        dataset,
        target,
        source_folder=folder,
    )


def rebuild_from_raw(
    folder: Path,
    *,
    actual_path: Path | None = None,
    report_path: Path | None = None,
    skip_invalid: bool = True,
) -> ReportSession:
    """Rebuild a report from raw brackets and save it, discarding saved edits."""

    session = build_from_raw(
        folder,
        actual_path=actual_path,
        report_path=report_path,
        skip_invalid=skip_invalid,
    )
    session.save()
    return session
