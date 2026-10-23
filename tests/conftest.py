"""Shared fixtures for the bracket scorekeeper tests."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SAMPLE_CSV = PROJECT_ROOT / "sample_data" / "2026_basketball_bracket_MAT.csv"


def read_rows(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [list(row) for row in csv.reader(handle)]


def write_rows(path: Path, rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle, lineterminator="\n").writerows(rows)


def make_sample(
    destination: Path,
    *,
    name: str = "MATT",
    champion: str | None = None,
    clear_cells: tuple[tuple[int, int], ...] = (),
    set_cells: tuple[tuple[int, int, str], ...] = (),
) -> Path:
    """Copy the full sample bracket to a CSV, optionally changing it.

    Rows and columns are 1-based, matching Excel coordinates.
    """

    rows = read_rows(SAMPLE_CSV)
    rows[1][0] = f"NAME:______________{name}_______________"
    if champion is not None:
        rows[42][8] = champion  # I43
    for row, col in clear_cells:
        rows[row - 1][col - 1] = ""
    for row, col, value in set_cells:
        rows[row - 1][col - 1] = value
    write_rows(destination, rows)
    return destination


def make_xlsx(
    source_csv: Path,
    destination: Path,
    *,
    merges: tuple[str, ...] = (
        "A1:AA1",
        "A2:Q2",
        "A6:C6",
        "O6:Q6",
        "A40:C40",
        "O40:Q40",
        "I44:J44",
    ),
) -> Path:
    """Convert the sample grid to an .xlsx, preserving merged-cell behavior."""

    rows = read_rows(source_csv)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Bracket"
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row, start=1):
            if value:
                worksheet.cell(row=row_index, column=col_index, value=value)
    for merged_range in merges:
        merged = worksheet[merged_range]
        top_left = merged[0][0]
        if top_left.value in (None, ""):
            # The sample text may place a label at the right edge of a merged
            # title block. Excel stores it at the top-left, so move it there.
            for row in merged:
                for cell in row:
                    if cell.value not in (None, ""):
                        top_left.value = cell.value
                        cell.value = None
                        break
                if top_left.value not in (None, ""):
                    break
        worksheet.merge_cells(merged_range)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()
    return destination


@pytest.fixture
def sample_csv() -> Path:
    return SAMPLE_CSV


@pytest.fixture
def sample_factory():
    return make_sample


@pytest.fixture
def xlsx_factory():
    return make_xlsx
