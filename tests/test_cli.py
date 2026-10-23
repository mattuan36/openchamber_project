"""End-to-end CLI behavior and exit codes."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from scorekeeper import main


def _write_demo(tmp_path: Path, sample_factory) -> Path:
    sample_factory(tmp_path / "2026_results.csv", name="OFFICIAL")
    sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    sample_factory(
        tmp_path / "2026_bracket_CHRIS.csv", name="CHRIS", champion="Houston"
    )
    return tmp_path


def test_cli_writes_report(tmp_path: Path, sample_factory, capsys) -> None:
    input_dir = _write_demo(tmp_path, sample_factory)
    output = tmp_path / "report.xlsx"
    code = main(["--input", str(input_dir), "--output", str(output)])
    captured = capsys.readouterr()
    assert code == 0
    assert output.exists()
    assert "Report written" in captured.out
    assert "MATT" in captured.out
    assert "148" in captured.out
    assert "No actual-results" not in captured.err
    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook.sheetnames == ["Picks", "Standings"]
    finally:
        workbook.close()


def test_cli_validate_only_does_not_write(
    tmp_path: Path, sample_factory, capsys
) -> None:
    input_dir = _write_demo(tmp_path, sample_factory)
    output = tmp_path / "should_not_exist.xlsx"
    code = main(
        [
            "--input",
            str(input_dir),
            "--output",
            str(output),
            "--validate-only",
        ]
    )
    captured = capsys.readouterr()
    assert code == 0
    assert not output.exists()
    assert "Validation OK" in captured.out
    assert "R64=32" in captured.out
    assert "Maximum points: 148" in captured.out


def test_cli_actual_override_allows_a_non_results_filename(
    tmp_path: Path, sample_factory, capsys
) -> None:
    actual = sample_factory(tmp_path / "official.csv", name="OFFICIAL")
    sample_factory(tmp_path / "alice.csv", name="ALICE")
    output = tmp_path / "report.xlsx"

    code = main(["--input", str(tmp_path), "--output", str(output)])
    captured = capsys.readouterr()
    assert code == 2
    assert "No actual-results" in captured.err
    assert not output.exists()

    code = main(
        [
            "--input",
            str(tmp_path),
            "--actual",
            str(actual),
            "--output",
            str(output),
        ]
    )
    captured = capsys.readouterr()
    assert code == 0
    assert output.exists()
    assert "ALICE" in captured.out


def test_cli_duplicate_results_is_an_error(
    tmp_path: Path, sample_factory, capsys
) -> None:
    _write_demo(tmp_path, sample_factory)
    sample_factory(tmp_path / "2026_other_results.csv", name="OTHER")
    code = main(
        [
            "--input",
            str(tmp_path),
            "--output",
            str(tmp_path / "report.xlsx"),
        ]
    )
    captured = capsys.readouterr()
    assert code == 2
    assert "Multiple files" in captured.err


def test_cli_rejects_results_in_report_filename(
    tmp_path: Path, sample_factory, capsys
) -> None:
    _write_demo(tmp_path, sample_factory)
    code = main(
        [
            "--input",
            str(tmp_path),
            "--output",
            str(tmp_path / "results_report.xlsx"),
        ]
    )
    captured = capsys.readouterr()
    assert code == 2
    assert "must not contain 'results'" in captured.err
    assert not (tmp_path / "results_report.xlsx").exists()


def test_cli_skip_invalid_entrants(
    tmp_path: Path, sample_factory, capsys
) -> None:
    _write_demo(tmp_path, sample_factory)
    sample_factory(
        tmp_path / "2026_bracket_BAD.csv",
        name="BAD",
        clear_cells=((6, 1), (6, 17)),
    )
    output = tmp_path / "report.xlsx"

    code = main(["--input", str(tmp_path), "--output", str(output)])
    captured = capsys.readouterr()
    assert code == 2
    assert "Template mismatch" in captured.err
    assert not output.exists()

    code = main(
        [
            "--input",
            str(tmp_path),
            "--output",
            str(output),
            "--skip-invalid",
        ]
    )
    captured = capsys.readouterr()
    assert code == 0
    assert output.exists()
    assert "Skipped 2026_bracket_BAD.csv" in captured.err


def test_cli_refuses_output_equal_to_actual(
    tmp_path: Path, sample_factory, capsys
) -> None:
    actual = sample_factory(tmp_path / "official.csv", name="OFFICIAL")
    sample_factory(tmp_path / "alice.csv", name="ALICE")
    code = main(
        [
            "--input",
            str(tmp_path),
            "--actual",
            str(actual),
            "--output",
            str(actual),
        ]
    )
    captured = capsys.readouterr()
    assert code == 2
    assert "output path cannot also be" in captured.err
    assert actual.exists()


def test_cli_refuses_output_that_is_an_excluded_entrant(
    tmp_path: Path, sample_factory, capsys
) -> None:
    import hashlib

    sample_factory(tmp_path / "2026_results.csv", name="OFFICIAL")
    entrant = sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    before = hashlib.sha256(entrant.read_bytes()).hexdigest()
    code = main(
        [
            "--input",
            str(tmp_path),
            "--output",
            str(entrant),
        ]
    )
    captured = capsys.readouterr()
    assert code == 2
    assert "Refusing to overwrite" in captured.err
    assert hashlib.sha256(entrant.read_bytes()).hexdigest() == before


def test_cli_allows_no_entrants(
    tmp_path: Path, sample_factory, capsys
) -> None:
    sample_factory(tmp_path / "2026_results.csv", name="OFFICIAL")
    output = tmp_path / "report.xlsx"
    code = main(["--input", str(tmp_path), "--output", str(output)])
    captured = capsys.readouterr()
    assert code == 0
    assert output.exists()
    assert "Entrants: 0" in captured.out
    assert "No entrant brackets" in captured.err


def test_cli_default_output_is_bracket_report(
    tmp_path: Path, sample_factory, monkeypatch, capsys
) -> None:
    _write_demo(tmp_path, sample_factory)
    monkeypatch.chdir(tmp_path)
    code = main(["--input", str(tmp_path)])
    captured = capsys.readouterr()
    assert code == 0
    assert (tmp_path / "bracket_report.xlsx").exists()
    assert "bracket_report.xlsx" in captured.out


def test_cli_actual_template_error_is_fatal(
    tmp_path: Path, sample_factory, capsys
) -> None:
    sample_factory(
        tmp_path / "2026_results.csv",
        name="OFFICIAL",
        set_cells=((8, 3, "17"),),
    )
    sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    code = main(
        [
            "--input",
            str(tmp_path),
            "--output",
            str(tmp_path / "report.xlsx"),
        ]
    )
    captured = capsys.readouterr()
    assert code == 2
    assert "seeds" in captured.err
