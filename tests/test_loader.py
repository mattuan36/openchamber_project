"""Visual-grid parsing, merged cells, discovery, and validation tests."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from bracket.loader import (
    BracketError,
    _grid_from_worksheet,
    discover_files,
    load_dataset,
    parse_bracket,
    select_actual_file,
)
from bracket.template import CANONICAL_SLOTS


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_parse_sample_csv_reconstructs_all_63_picks(sample_csv: Path) -> None:
    bracket = parse_bracket(sample_csv)
    assert bracket.entrant == "MATT"
    assert bracket.filled == 63
    assert bracket.picks["CH"] == "Arizona"
    assert bracket.picks["F4-1"] == "Houston"
    assert bracket.picks["F4-2"] == "Arizona"
    assert bracket.picks["E8-E1"] == "Duke"
    assert bracket.picks["R64-W1"] == "Arizona"
    assert all("CSV preview" in warning for warning in bracket.warnings)


def test_parse_xlsx_with_merged_anchor_cells(
    sample_csv: Path, tmp_path: Path, xlsx_factory
) -> None:
    workbook_path = tmp_path / "2026_bracket_MATT.xlsx"
    xlsx_factory(sample_csv, workbook_path)
    bracket = parse_bracket(workbook_path)
    assert bracket.entrant == "MATT"
    assert bracket.filled == 63
    assert bracket.picks["CH"] == "Arizona"
    assert bracket.warnings == ()


def test_grid_resolves_values_inside_merged_ranges(tmp_path: Path) -> None:
    path = tmp_path / "merged.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bracket"
    sheet["D9"] = "Duke"
    sheet.merge_cells("D9:F10")
    workbook.save(path)
    workbook.close()

    opened = load_workbook(path, data_only=True)
    try:
        grid = _grid_from_worksheet(opened["Bracket"])
    finally:
        opened.close()
    assert grid.value(9, 4) == "Duke"
    assert grid.value(10, 5) == "Duke"
    assert grid.value(11, 4) == ""


def test_right_team_column_can_be_p_or_q(
    sample_csv: Path, tmp_path: Path
) -> None:
    from tests.conftest import read_rows, write_rows

    rows = read_rows(sample_csv)
    # Move the right-hand team labels from Q to P and put the record in Q.
    # Detection must prefer the letters in P over the digits in Q.
    for row_index in range(1, 72):
        if row_index in (6, 40):
            continue
        value = rows[row_index - 1][16]
        record = rows[row_index - 1][17]
        if value and rows[row_index - 1][14]:
            rows[row_index - 1][15] = value
            rows[row_index - 1][16] = record
            rows[row_index - 1][17] = ""
    path = tmp_path / "2026_bracket_LEFT.xlsx"
    converted = tmp_path / "2026_bracket_LEFT.csv"
    write_rows(converted, rows)
    from tests.conftest import make_xlsx

    make_xlsx(converted, path)
    bracket = parse_bracket(path)
    assert bracket.filled == 63
    assert bracket.picks["CH"] == "Arizona"

    # The mirrored layout: record in P, team in Q.
    rows = read_rows(sample_csv)
    for row_index in range(1, 72):
        if row_index in (6, 40):
            continue
        record = rows[row_index - 1][17]
        value = rows[row_index - 1][16]
        if value and rows[row_index - 1][14]:
            rows[row_index - 1][15] = record
            rows[row_index - 1][16] = value
            rows[row_index - 1][17] = ""
    path = tmp_path / "2026_bracket_RIGHT.xlsx"
    converted = tmp_path / "2026_bracket_RIGHT.csv"
    write_rows(converted, rows)
    make_xlsx(converted, path)
    bracket = parse_bracket(path)
    assert bracket.filled == 63
    assert bracket.picks["CH"] == "Arizona"


def _make_three_files(tmp_path: Path, sample_factory) -> tuple[Path, Path, Path]:
    actual = sample_factory(tmp_path / "2026_results.csv", name="OFFICIAL")
    matt = sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    chris = sample_factory(
        tmp_path / "2026_bracket_CHRIS.csv", name="CHRIS", champion="Houston"
    )
    return actual, matt, chris


def test_discovery_excludes_output_and_ignores_non_brackets(
    tmp_path: Path, sample_factory
) -> None:
    _make_three_files(tmp_path, sample_factory)
    (tmp_path / "notes.txt").write_text("not a bracket", encoding="utf-8")
    (tmp_path / "~$2026_bracket_MATT.xlsx").write_text("lock", encoding="utf-8")
    output = tmp_path / "bracket_report.xlsx"
    output.write_text("placeholder", encoding="utf-8")

    files = discover_files(tmp_path, output)
    assert [path.name for path in files] == [
        "2026_bracket_CHRIS.csv",
        "2026_bracket_MATT.csv",
        "2026_results.csv",
    ]
    assert select_actual_file(files, None, output).name == "2026_results.csv"


def test_actual_detection_is_case_insensitive_and_overrideable(
    tmp_path: Path, sample_factory
) -> None:
    actual = sample_factory(tmp_path / "OFFICIAL_RESULTS.CSV", name="OFFICIAL")
    files = discover_files(tmp_path)
    assert select_actual_file(files, None, None) == actual.resolve()

    other = sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    assert select_actual_file(files, other, None) == other.resolve()


def test_multiple_or_missing_actual_files_are_clear_errors(
    tmp_path: Path, sample_factory
) -> None:
    sample_factory(tmp_path / "2026_results.csv", name="OFFICIAL")
    sample_factory(tmp_path / "2026_more_results.csv", name="OFFICIAL2")
    files = discover_files(tmp_path)
    with pytest.raises(BracketError, match="Multiple files"):
        select_actual_file(files, None, None)

    empty = tmp_path / "entrants"
    empty.mkdir()
    sample_factory(empty / "2026_bracket_MATT.csv", name="MATT")
    with pytest.raises(BracketError, match="No actual-results"):
        select_actual_file(discover_files(empty), None, None)


def test_actual_cannot_be_the_output(
    tmp_path: Path, sample_factory
) -> None:
    actual = sample_factory(tmp_path / "2026_results.csv", name="OFFICIAL")
    with pytest.raises(BracketError, match="output path"):
        select_actual_file([actual], actual, actual)


def test_load_dataset_and_duplicate_name_warning(
    tmp_path: Path, sample_factory
) -> None:
    _make_three_files(tmp_path, sample_factory)
    dataset = load_dataset(tmp_path, output_path=tmp_path / "bracket_report.xlsx")
    assert dataset.actual.entrant == "OFFICIAL"
    assert dataset.actual.filled == 63
    assert [bracket.entrant for bracket in dataset.entrants] == ["CHRIS", "MATT"]
    assert dataset.pending == 0
    assert len(dataset.files) == 3
    assert dataset.warnings == ()

    duplicate = sample_factory(tmp_path / "2026_bracket_MATT2.csv", name="MATT")
    dataset = load_dataset(tmp_path, output_path=tmp_path / "bracket_report.xlsx")
    assert duplicate.exists()
    assert any("Duplicate entrant name" in warning for warning in dataset.warnings)


def test_invalid_entrant_can_be_skipped(
    tmp_path: Path, sample_factory
) -> None:
    sample_factory(tmp_path / "2026_results.csv", name="OFFICIAL")
    sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    sample_factory(
        tmp_path / "2026_bracket_BAD.csv",
        name="BAD",
        clear_cells=((6, 1), (6, 17)),
    )
    with pytest.raises(BracketError, match="Template mismatch"):
        load_dataset(tmp_path, output_path=tmp_path / "report.xlsx")

    dataset = load_dataset(
        tmp_path,
        output_path=tmp_path / "report.xlsx",
        skip_invalid=True,
    )
    assert [bracket.entrant for bracket in dataset.entrants] == ["MATT"]
    assert len(dataset.skipped) == 1
    assert dataset.skipped[0][0].name == "2026_bracket_BAD.csv"


def test_actual_template_errors_are_never_skipped(
    tmp_path: Path, sample_factory
) -> None:
    sample_factory(
        tmp_path / "2026_results.csv",
        name="OFFICIAL",
        set_cells=((8, 3, "17"),),
    )
    sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    with pytest.raises(BracketError, match="seeds"):
        load_dataset(
            tmp_path,
            output_path=tmp_path / "report.xlsx",
            skip_invalid=True,
        )

    # Missing champion label is also a fatal template mismatch.
    sample_factory(
        tmp_path / "2026_results.csv",
        name="OFFICIAL",
        clear_cells=((44, 9),),
    )
    with pytest.raises(BracketError, match="Template mismatch"):
        load_dataset(
            tmp_path,
            output_path=tmp_path / "report.xlsx",
            skip_invalid=True,
        )


def test_actual_with_no_picks_is_rejected(
    tmp_path: Path, sample_factory
) -> None:
    sample_factory(
        tmp_path / "2026_results.csv",
        name="OFFICIAL",
        clear_cells=tuple((slot.row, slot.col) for slot in CANONICAL_SLOTS),
    )
    sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    with pytest.raises(BracketError, match="no picks"):
        load_dataset(tmp_path, output_path=tmp_path / "report.xlsx")


def test_output_is_excluded_from_discovery(
    tmp_path: Path, sample_factory
) -> None:
    _make_three_files(tmp_path, sample_factory)
    output = tmp_path / "bracket_report.xlsx"
    output.write_text("old report", encoding="utf-8")
    dataset = load_dataset(tmp_path, output_path=output)
    assert all(path.name != "bracket_report.xlsx" for path in dataset.files)
    assert len(dataset.entrants) == 2


def test_reading_inputs_does_not_modify_them(
    tmp_path: Path, sample_factory
) -> None:
    paths = _make_three_files(tmp_path, sample_factory)
    before = {path: _digest(path) for path in paths}
    load_dataset(tmp_path, output_path=tmp_path / "bracket_report.xlsx")
    after = {path: _digest(path) for path in paths}
    assert before == after
