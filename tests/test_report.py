"""Report workbook layout, formatting, and safe-write tests."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bracket.loader import load_dataset
from bracket.report import ReportError, build_report, write_report
from bracket.scoring import score_all
from bracket.template import SLOT_IDS


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _champion_column() -> int:
    return 2 + len(SLOT_IDS)


def _f4_left_column() -> int:
    return 3 + SLOT_IDS.index("F4-1")


def _report_dataset(tmp_path: Path, sample_factory):
    sample_factory(tmp_path / "2026_results.csv", name="OFFICIAL")
    sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    sample_factory(
        tmp_path / "2026_bracket_CHRIS.csv", name="CHRIS", champion="Houston"
    )
    dataset = load_dataset(tmp_path, output_path=tmp_path / "bracket_report.xlsx")
    return dataset, score_all(dataset.actual, dataset.entrants)


def test_report_has_five_sheets_and_63_game_columns(
    tmp_path: Path, sample_factory
) -> None:
    dataset, scores = _report_dataset(tmp_path, sample_factory)
    workbook = build_report(dataset, scores)
    try:
        assert workbook.sheetnames == [
            "Picks",
            "Standings",
            "State",
            "Teams",
            "Corrections",
        ]
        assert all(
            workbook[name].sheet_state == "visible" for name in workbook.sheetnames
        )
        picks = workbook["Picks"]
        assert picks.max_column == 2 + len(SLOT_IDS)
        assert [picks.cell(row=1, column=col).value for col in (1, 2)] == [
            "Source File",
            "Entrant",
        ]
        labels = [
            picks.cell(row=1, column=col).value
            for col in range(3, picks.max_column + 1)
        ]
        assert len(labels) == 63
        assert labels[0] == "R64 East 1"
        assert labels[7] == "R64 East 8"
        assert labels[8] == "R64 West 1"
        assert labels[31] == "R64 Midwest 8"
        assert labels[32] == "R32 East 1"
        assert labels[-3:] == ["F4 Game 1", "F4 Game 2", "Championship"]
        assert picks.freeze_panes == "C2"
        assert picks.auto_filter.ref == f"A1:BM{picks.max_row}"
    finally:
        workbook.close()


def test_picks_rows_are_actual_then_ranked_entrants(
    tmp_path: Path, sample_factory
) -> None:
    dataset, scores = _report_dataset(tmp_path, sample_factory)
    workbook = build_report(dataset, scores)
    try:
        picks = workbook["Picks"]
        assert picks["A2"].value == "2026_results.csv"
        assert picks["B2"].value == "ACTUAL — OFFICIAL"
        assert picks["B3"].value == "MATT"
        assert picks["B4"].value == "CHRIS"
        assert picks["C2"].fill.fgColor.rgb == "001F4E78"
        assert picks["C3"].fill.fgColor.rgb == "00C6EFCE"
        assert picks["C4"].fill.fgColor.rgb == "00C6EFCE"
        champion_col = _champion_column()
        assert picks.cell(row=3, column=champion_col).value == "Arizona"
        assert picks.cell(row=3, column=champion_col).fill.fgColor.rgb == "00C6EFCE"
        assert picks.cell(row=4, column=champion_col).value == "Houston"
        assert picks.cell(row=4, column=champion_col).fill.fgColor.rgb == "00FFC7CE"
    finally:
        workbook.close()


def test_pending_cells_are_gray_and_not_counted_as_wrong(
    tmp_path: Path, sample_factory
) -> None:
    sample_factory(
        tmp_path / "2026_results.csv",
        name="OFFICIAL",
        clear_cells=((43, 9), (34, 8)),
    )
    sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    dataset = load_dataset(tmp_path, output_path=tmp_path / "bracket_report.xlsx")
    scores = score_all(dataset.actual, dataset.entrants)
    assert scores[0].pending == 2
    assert scores[0].correct == 61
    assert scores[0].total == 128
    assert scores[0].accuracy == 1.0

    workbook = build_report(dataset, scores)
    try:
        picks = workbook["Picks"]
        champion_col = _champion_column()
        f4_col = _f4_left_column()
        assert picks.cell(row=2, column=champion_col).fill.fgColor.rgb == "00E7E6E6"
        assert picks.cell(row=3, column=champion_col).fill.fgColor.rgb == "00E7E6E6"
        assert picks.cell(row=3, column=champion_col).value == "Arizona"
        assert picks.cell(row=2, column=f4_col).fill.fgColor.rgb == "00E7E6E6"
        assert picks.cell(row=3, column=f4_col).value == "Houston"

        standings = workbook["Standings"]
        assert standings["A2"].value == 1
        assert standings["C2"].value == "2026_bracket_MATT.csv"
        assert standings["J2"].value == 128
        assert standings["K2"].value == 61
        assert standings["L2"].value == 2
        assert standings["M2"].value == 1.0
        assert standings["M2"].number_format == "0.0%"
    finally:
        workbook.close()


def test_standings_contains_points_rank_and_notes(
    tmp_path: Path, sample_factory
) -> None:
    dataset, scores = _report_dataset(tmp_path, sample_factory)
    workbook = build_report(dataset, scores)
    try:
        standings = workbook["Standings"]
        assert [standings.cell(row=1, column=col).value for col in range(1, 14)] == [
            "Rank",
            "Entrant",
            "Source File",
            "R64 (×1)",
            "R32 (×2)",
            "S16 (×4)",
            "E8 (×6)",
            "F4 (×8)",
            "Championship (×12)",
            "Total",
            "Correct Picks",
            "Pending Games",
            "Accuracy",
        ]
        assert standings["A2"].value == 1
        assert standings["B2"].value == "MATT"
        assert standings["D2"].value == 32
        assert standings["E2"].value == 32
        assert standings["F2"].value == 32
        assert standings["G2"].value == 24
        assert standings["H2"].value == 16
        assert standings["I2"].value == 12
        assert standings["J2"].value == 148
        assert standings["K2"].value == 63
        assert standings["L2"].value == 0
        assert standings["M2"].value == 1.0
        assert standings["A3"].value == 2
        assert standings["B3"].value == "CHRIS"
        assert standings["J3"].value == 136
        assert standings["M3"].value == pytest.approx(62 / 63)
        assert standings.freeze_panes == "C2"
        assert standings.auto_filter.ref == "A1:M3"
        assert standings.merged_cells.ranges
        assert any("Scoring:" in str(standings.cell(row=row, column=1).value) for row in range(5, 8))
        assert any("pending" in str(standings.cell(row=row, column=1).value).lower() for row in range(5, 8))
        assert any("Ties are broken" in str(standings.cell(row=row, column=1).value) for row in range(5, 8))
        assert list(standings.conditional_formatting)
    finally:
        workbook.close()


def test_write_report_is_atomic_and_preserves_inputs(
    tmp_path: Path, sample_factory
) -> None:
    dataset, scores = _report_dataset(tmp_path, sample_factory)
    input_paths = [dataset.actual.path, *[bracket.path for bracket in dataset.entrants]]
    before = {path: _digest(path) for path in input_paths}
    output = tmp_path / "out" / "bracket_report.xlsx"
    result = write_report(dataset, scores, output)
    assert result == output.resolve()
    assert output.exists()
    assert {path: _digest(path) for path in input_paths} == before
    assert not list(output.parent.glob(".bracket_report_*.xlsx"))

    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook.sheetnames == [
            "Picks",
            "Standings",
            "State",
            "Teams",
            "Corrections",
        ]
    finally:
        workbook.close()


def test_write_report_refuses_to_overwrite_an_input(
    tmp_path: Path, sample_factory
) -> None:
    dataset, scores = _report_dataset(tmp_path, sample_factory)
    before = _digest(dataset.actual.path)
    with pytest.raises(ReportError, match="Refusing to overwrite"):
        write_report(dataset, scores, dataset.actual.path)
    assert _digest(dataset.actual.path) == before


def test_write_report_refuses_an_excluded_entrant_as_output(
    tmp_path: Path, sample_factory
) -> None:
    sample_factory(tmp_path / "2026_results.csv", name="OFFICIAL")
    entrant = sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    dataset = load_dataset(tmp_path, output_path=entrant)
    scores = score_all(dataset.actual, dataset.entrants)
    assert entrant.resolve() not in dataset.files
    before = _digest(entrant)
    with pytest.raises(ReportError, match="Refusing to overwrite"):
        write_report(dataset, scores, entrant)
    assert _digest(entrant) == before


def test_write_report_can_replace_a_previous_report(
    tmp_path: Path, sample_factory
) -> None:
    dataset, scores = _report_dataset(tmp_path, sample_factory)
    output = tmp_path / "bracket_report.xlsx"
    write_report(dataset, scores, output)
    write_report(dataset, scores, output)
    assert output.exists()


def test_old_generated_report_is_not_discovered_as_an_entrant(
    tmp_path: Path, sample_factory
) -> None:
    dataset, scores = _report_dataset(tmp_path, sample_factory)
    old_report = tmp_path / "old_report.xlsx"
    write_report(dataset, scores, old_report)
    reloaded = load_dataset(
        tmp_path, output_path=tmp_path / "new_report.xlsx"
    )
    assert [bracket.entrant for bracket in reloaded.entrants] == ["CHRIS", "MATT"]
    assert all(path.name != "old_report.xlsx" for path in reloaded.files)


def test_write_report_with_no_entrants(
    tmp_path: Path, sample_factory
) -> None:
    sample_factory(tmp_path / "2026_results.csv", name="OFFICIAL")
    dataset = load_dataset(tmp_path, output_path=tmp_path / "bracket_report.xlsx")
    scores = score_all(dataset.actual, dataset.entrants)
    workbook = build_report(dataset, scores)
    try:
        assert workbook["Picks"].max_row == 2
        standings = workbook["Standings"]
        assert "No entrant brackets" in standings["A2"].value
    finally:
        workbook.close()

    output = tmp_path / "final.xlsx"
    write_report(dataset, scores, output)
    assert output.exists()
