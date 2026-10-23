"""Five-sheet report persistence, correction, and local-folder workflow tests."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bracket.loader import BracketError, is_generated_report, load_dataset
from bracket.report import ReportError
from bracket.scoring import available_picks
from bracket.store import (
    build_from_raw,
    load_report,
    rebuild_from_raw,
    scan_folder,
)
from bracket.template import SLOT_IDS


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _demo(tmp_path: Path, sample_factory) -> Path:
    sample_factory(tmp_path / "2026_results.csv", name="OFFICIAL")
    sample_factory(tmp_path / "2026_bracket_MATT.csv", name="MATT")
    sample_factory(
        tmp_path / "2026_bracket_CHRIS.csv", name="CHRIS", champion="Houston"
    )
    return tmp_path


def _session(tmp_path: Path, sample_factory):
    _demo(tmp_path, sample_factory)
    return build_from_raw(tmp_path, report_path=tmp_path / "bracket_report.xlsx")


def _index_of(session, entrant: str) -> int:
    return next(
        index
        for index, bracket in enumerate(session.entrants)
        if bracket.entrant == entrant
    )


def test_report_round_trip_preserves_picks_corrections_and_scores(
    tmp_path: Path, sample_factory
) -> None:
    session = _session(tmp_path, sample_factory)
    target = _index_of(session, "MATT")
    session, _entry = session.with_corrected_entrant(
        target,
        "R64-E1",
        "Siena",
        "Misread the bracket",
        timestamp="2026-03-20T12:00:00+00:00",
    )
    session.save()

    reloaded = load_report(tmp_path / "bracket_report.xlsx")
    assert reloaded.actual.entrant == "OFFICIAL"
    assert reloaded.actual.source_name == "2026_results.csv"
    assert reloaded.actual.picks == session.actual.picks
    assert {
        bracket.entrant: bracket.picks for bracket in reloaded.entrants
    } == {bracket.entrant: bracket.picks for bracket in session.entrants}
    assert reloaded.corrections == session.corrections
    assert [
        (score.entrant, score.total, score.rank) for score in reloaded.scores
    ] == [(score.entrant, score.total, score.rank) for score in session.scores]
    assert len(reloaded.team_pairs) == 32
    assert reloaded.team_pairs["R64-E1"] == ("Duke", "Siena")
    assert reloaded.team_seeds["R64-E1"] == (1, 16)
    assert reloaded.team_seeds["R64-W5"] == (6, 11, 11)
    assert reloaded.team_seeds_complete
    assert reloaded.team_seed_map["duke"] == 1
    assert reloaded.team_seed_map["north carolina state"] == 11
    assert reloaded.state["Report schema version"] == "3"
    assert reloaded.state["Scoring rules version"] == "2026.1"


def test_state_teams_and_corrections_sheet_contents(
    tmp_path: Path, sample_factory
) -> None:
    session = _session(tmp_path, sample_factory)
    session, _entry = session.with_corrected_entrant(
        _index_of(session, "MATT"),
        "CH",
        "Houston",
        "Championship pick was recorded wrong",
    )
    session.save()

    workbook = load_workbook(tmp_path / "bracket_report.xlsx", data_only=True)
    try:
        assert workbook.sheetnames == [
            "Picks",
            "Standings",
            "State",
            "Teams",
            "Corrections",
        ]

        state = workbook["State"]
        keys = [state.cell(row=row, column=1).value for row in range(2, 9)]
        values = [str(state.cell(row=row, column=2).value or "") for row in range(2, 9)]
        assert keys == [
            "Report schema version",
            "Created (UTC)",
            "Updated (UTC)",
            "Source folder",
            "Scoring rules version",
            "App version",
            "Report format",
        ]
        assert all(values)
        assert values[1].endswith("+00:00")
        assert values[2].endswith("+00:00")
        assert "MATT" not in " ".join(values)
        assert "CHRIS" not in " ".join(values)
        assert "2026_bracket_MATT.csv" not in values
        assert "2026_results.csv" not in values

        teams = workbook["Teams"]
        assert [teams.cell(row=1, column=col).value for col in range(1, 7)] == [
            "Slot",
            "Region",
            "Seed 1",
            "Team 1",
            "Seed 2",
            "Team 2",
        ]
        rows = [
            [teams.cell(row=row, column=col).value for col in range(1, 7)]
            for row in range(2, teams.max_row + 1)
        ]
        assert len(rows) == 32
        assert len({row[0] for row in rows}) == 32
        by_slot = {row[0]: row for row in rows}
        assert by_slot["R64-E1"] == ["R64-E1", "East", 1, "Duke", 16, "Siena"]
        assert by_slot["R64-W5"] == [
            "R64-W5",
            "West",
            6,
            "BYU",
            11,
            "Texas / NC State",
        ]
        assert all(row[0].startswith("R64-") for row in rows)

        corrections = workbook["Corrections"]
        assert [
            corrections.cell(row=1, column=col).value for col in range(1, 8)
        ] == [
            "Timestamp",
            "Entrant",
            "Source File",
            "Slot",
            "Old Pick",
            "New Pick",
            "Reason",
        ]
        assert corrections["B2"].value == "MATT"
        assert corrections["D2"].value == "CH"
        assert corrections["E2"].value == "Arizona"
        assert corrections["F2"].value == "Houston"
        assert corrections["G2"].value == "Championship pick was recorded wrong"
    finally:
        workbook.close()


def test_legacy_five_sheet_report_without_seeds_loads_names_only(
    tmp_path: Path, sample_factory
) -> None:
    session = _session(tmp_path, sample_factory)
    report = session.save()

    workbook = load_workbook(report)
    try:
        teams = workbook["Teams"]
        teams["C1"] = "Team 1"
        teams["D1"] = "Team 2"
        teams["E1"] = None
        teams["F1"] = None
        for row in range(2, teams.max_row + 1):
            teams.cell(row=row, column=3).value = teams.cell(row=row, column=4).value
            teams.cell(row=row, column=4).value = teams.cell(row=row, column=6).value
            teams.cell(row=row, column=5).value = None
            teams.cell(row=row, column=6).value = None
        workbook.save(report)
    finally:
        workbook.close()

    reloaded = load_report(report)
    assert reloaded.team_seeds == {}
    assert not reloaded.team_seeds_complete
    assert reloaded.teams_complete
    assert reloaded.team_pairs["R64-E1"] == ("Duke", "Siena")
    assert reloaded.team_pairs["R64-W5"] == ("BYU", "Texas", "NC State")
    assert reloaded.team_seed_map == {}


def test_report_loads_by_header_name_after_column_reorder(
    tmp_path: Path, sample_factory
) -> None:
    session = _session(tmp_path, sample_factory)
    report = session.save()
    expected_actual = dict(session.actual.picks)
    expected_entrants = {bracket.entrant: dict(bracket.picks) for bracket in session.entrants}

    workbook = load_workbook(report)
    try:
        sheet = workbook["Picks"]
        rows = list(sheet.iter_rows(values_only=True))
        order = [0, 1, *range(len(rows[0]) - 1, 1, -1)]
        reordered = [[row[index] for index in order] for row in rows]
        for row_index in range(1, sheet.max_row + 1):
            for column_index in range(1, sheet.max_column + 1):
                sheet.cell(row=row_index, column=column_index).value = None
        for row_index, row in enumerate(reordered, start=1):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row=row_index, column=column_index).value = value
        workbook.save(report)
    finally:
        workbook.close()

    reloaded = load_report(report)
    assert reloaded.actual.picks == expected_actual
    assert {
        bracket.entrant: bracket.picks for bracket in reloaded.entrants
    } == expected_entrants


def test_actual_cascade_pending_and_available_picks(
    tmp_path: Path, sample_factory
) -> None:
    session = _session(tmp_path, sample_factory)
    assert session.actual.picks["R32-E1"] == "Duke"
    assert available_picks(
        "R32-E1", session.team_pairs, session.actual.picks
    ) == ["Duke", "Ohio State"]

    changed, cleared = session.with_actual_pick("R64-E1", "Siena")
    assert cleared == ["R32-E1", "S16-E1", "E8-E1"]
    assert changed.actual.picks["R64-E1"] == "Siena"
    assert changed.actual.picks["R32-E1"] == ""
    assert changed.actual.picks["S16-E1"] == ""
    assert changed.actual.picks["E8-E1"] == ""
    assert changed.pending == 3
    assert available_picks(
        "R32-E1", changed.team_pairs, changed.actual.picks
    ) == ["Siena", "Ohio State"]
    assert changed.scores[0].pending == 3
    assert changed.scores[0].total == 135

    # A downstream pick is not restored automatically when a feeder returns.
    restored, cleared_again = changed.with_actual_pick("R64-E1", "Duke")
    assert cleared_again == []
    assert restored.actual.picks["R32-E1"] == ""


def test_entrant_correction_requires_reason_and_is_logged(
    tmp_path: Path, sample_factory
) -> None:
    session = _session(tmp_path, sample_factory)
    target = _index_of(session, "MATT")
    with pytest.raises(ReportError, match="reason"):
        session.with_corrected_entrant(target, "CH", "Houston", "   ")
    with pytest.raises(ReportError, match="feeder winners for CH"):
        session.with_corrected_entrant(target, "CH", "Siena", "Wrong team")

    corrected, entry = session.with_corrected_entrant(
        target, "CH", "Houston", "Pick was transcribed from the wrong row"
    )
    assert entry["Old Pick"] == "Arizona"
    assert entry["New Pick"] == "Houston"
    assert entry["Entrant"] == "MATT"
    assert corrected.corrections == [entry]
    assert corrected.entrants[target].picks["CH"] == "Houston"
    matts = next(score for score in corrected.scores if score.entrant == "MATT")
    assert matts.total == 136

    corrected.save()
    reloaded = load_report(corrected.report_path)
    assert reloaded.corrections == [entry]
    matts = next(
        bracket for bracket in reloaded.entrants if bracket.entrant == "MATT"
    )
    assert matts.picks["CH"] == "Houston"


def test_save_refuses_raw_bracket_and_allows_overwriting_report(
    tmp_path: Path, sample_factory, xlsx_factory, sample_csv
) -> None:
    raw_xlsx = xlsx_factory(sample_csv, tmp_path / "2026_bracket_RAW.xlsx")
    _demo(tmp_path, sample_factory)
    session = build_from_raw(tmp_path, report_path=tmp_path / "bracket_report.xlsx")
    before = _digest(raw_xlsx)

    with pytest.raises(ReportError, match="Refusing to overwrite"):
        session.save(raw_xlsx)
    assert _digest(raw_xlsx) == before

    with pytest.raises(ReportError, match="must not contain 'results'"):
        session.save(tmp_path / "results_report.xlsx")

    session.save()
    assert (tmp_path / "bracket_report.xlsx").exists()
    assert load_report(tmp_path / "bracket_report.xlsx").actual.picks == session.actual.picks
    session.save()
    assert load_report(tmp_path / "bracket_report.xlsx").actual.picks == session.actual.picks


def test_locked_report_keeps_unsaved_state(
    tmp_path: Path, sample_factory, monkeypatch
) -> None:
    session = _session(tmp_path, sample_factory)
    report = session.save()
    before = _digest(report)

    def locked_replace(source, destination):
        raise PermissionError("file is open in Excel")

    monkeypatch.setattr("bracket.report.os.replace", locked_replace)
    updated, _cleared = session.with_actual_pick("R64-E1", "Siena")
    with pytest.raises(ReportError, match="Close it in Excel"):
        updated.save()

    assert _digest(report) == before
    assert updated.report_path == report
    assert updated.actual.picks["R64-E1"] == "Siena"
    assert not list(report.parent.glob(".bracket_report_*.xlsx"))


def test_legacy_two_sheet_report_still_loads(
    tmp_path: Path, sample_factory
) -> None:
    session = _session(tmp_path, sample_factory)
    report = session.save()
    workbook = load_workbook(report)
    try:
        for name in ("State", "Teams", "Corrections"):
            del workbook[name]
        assert workbook.sheetnames == ["Picks", "Standings"]
        workbook.save(report)
    finally:
        workbook.close()

    assert is_generated_report(report)
    reloaded = load_report(report)
    assert reloaded.actual.picks == session.actual.picks
    assert [bracket.entrant for bracket in reloaded.entrants] == ["MATT", "CHRIS"]
    assert reloaded.team_pairs == {}
    assert not reloaded.teams_complete
    assert any("Legacy report" in warning for warning in reloaded.warnings)
    assert reloaded.state["Report format"] == "Legacy two-sheet report"
    assert reloaded.corrections == []
    with pytest.raises(ReportError, match="Rebuild it from the original bracket files"):
        reloaded.save()


def test_local_folder_scan_build_and_rebuild(
    tmp_path: Path, sample_factory
) -> None:
    _demo(tmp_path, sample_factory)
    # Excel owner files and nested files must not become raw entrants.
    (tmp_path / "~$2026_bracket_MATT.xlsx").write_bytes(b"owner file")
    nested = tmp_path / "nested"
    nested.mkdir()
    (nested / "2026_bracket_NESTED.csv").write_text("not scanned", encoding="utf-8")

    scan = scan_folder(tmp_path)
    assert {path.name for path in scan.raw_files} == {
        "2026_bracket_CHRIS.csv",
        "2026_bracket_MATT.csv",
        "2026_results.csv",
    }
    assert scan.reports == []
    assert [path.name for path in scan.actual_candidates] == ["2026_results.csv"]

    session = build_from_raw(tmp_path)
    report = session.save()
    assert report == (tmp_path / "bracket_report.xlsx").resolve()

    scan = scan_folder(tmp_path)
    assert [path.name for path in scan.reports] == ["bracket_report.xlsx"]
    assert len(scan.raw_files) == 3

    edited, _cleared = session.with_actual_pick("R64-E1", "Siena")
    edited.save()
    assert load_report(report).actual.picks["R64-E1"] == "Siena"

    rebuilt = rebuild_from_raw(tmp_path, report_path=report)
    assert rebuilt.actual.picks["R64-E1"] == "Duke"
    assert load_report(report).actual.picks["R64-E1"] == "Duke"
    assert rebuilt.corrections == []


def test_folder_with_only_a_report_can_still_continue(
    tmp_path: Path, sample_factory
) -> None:
    session = _session(tmp_path, sample_factory)
    report = session.save()
    folder = tmp_path / "only_report"
    folder.mkdir()
    (folder / report.name).write_bytes(report.read_bytes())

    scan = scan_folder(folder)
    assert scan.raw_files == []
    assert [path.name for path in scan.reports] == ["bracket_report.xlsx"]
    with pytest.raises(BracketError):
        build_from_raw(folder)


def test_dataset_warns_when_entrant_first_round_teams_disagree(
    tmp_path: Path, sample_factory
) -> None:
    _demo(tmp_path, sample_factory)
    sample_factory(
        tmp_path / "2026_bracket_OTHER.csv",
        name="OTHER",
        set_cells=((8, 1, "Different Team"),),
    )
    dataset = load_dataset(tmp_path, output_path=tmp_path / "bracket_report.xlsx")
    assert any("first-round teams" in warning for warning in dataset.warnings)


def test_picks_has_one_column_for_every_canonical_slot(
    tmp_path: Path, sample_factory
) -> None:
    session = _session(tmp_path, sample_factory)
    report = session.save()
    workbook = load_workbook(report, read_only=True)
    try:
        assert workbook["Picks"].max_column == 2 + len(SLOT_IDS)
    finally:
        workbook.close()


def test_stored_standings_are_ignored_and_recalculated(
    tmp_path: Path, sample_factory
) -> None:
    session = _session(tmp_path, sample_factory)
    report = session.save()

    workbook = load_workbook(report)
    try:
        standings = workbook["Standings"]
        standings["A2"] = 99
        standings["B2"] = "TAMPERED"
        standings["J2"] = 1
        standings["K2"] = 0
        standings["M2"] = 0.01
        workbook.save(report)
    finally:
        workbook.close()

    reloaded = load_report(report)
    assert [
        (score.entrant, score.total, score.correct, score.rank)
        for score in reloaded.scores
    ] == [("MATT", 148, 63, 1), ("CHRIS", 136, 62, 2)]
